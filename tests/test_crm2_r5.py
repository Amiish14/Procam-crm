"""
CRM 2.0 R5 — the forwarding machinery, asked to answer for every
vertical.

The thing these tests are really guarding is that there is still one
of everything. It would have been quicker to copy ff_rates into
transport_rates and ff_charges into transport_charges, and within a
quarter the two would disagree about what a margin is. So several
tests here assert the absence of a second copy rather than the
presence of a feature.
"""
import os
import sys
import tempfile
from datetime import date

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

os.environ['URL_PREFIX'] = ''
os.environ.setdefault('SECRET_KEY', 'r5-secret')
os.environ.setdefault('ADMIN_INITIAL_PASSWORD', 'R5Test-12345')
os.environ.setdefault(
    'DATABASE_URL',
    'sqlite:///' + os.path.join(tempfile.mkdtemp(), 'r5.db'))

import importlib                                          # noqa: E402
_main = importlib.import_module('app')
importlib.import_module('app.models')
flask_app = _main.app
db = _main.db
Lead = _main.Lead

from app.services import (ff_charges, job_handover,        # noqa: E402
                          quote_build, quote_send, quote_terms)


# ── charge groups per vertical ───────────────────────────────────────
def test_each_vertical_offers_the_groups_it_actually_uses():
    assert 'export_customs' in ff_charges.groups_for(
        'Freight Forwarding')
    assert 'export_customs' not in ff_charges.groups_for(
        'General Transport')
    assert 'escort_permits' in ff_charges.groups_for('Heavy Transport')
    assert 'storage' in ff_charges.groups_for('Warehousing')


def test_other_is_on_every_list():
    """The charge nobody anticipated is the one that costs money."""
    for vertical in ff_charges.VERTICAL_GROUPS:
        assert 'other' in ff_charges.groups_for(vertical)


def test_an_unknown_vertical_falls_back_rather_than_crashing():
    assert ff_charges.groups_for('Something New') == \
        ff_charges.groups_for('Freight Forwarding')
    assert ff_charges.groups_for(None)


def test_a_lorry_has_no_incoterm():
    """Asking for one would be a field nobody can fill, so everything
    on the quote is the client's."""
    assert not ff_charges.uses_incoterms('General Transport')
    billable = ff_charges.billable_groups(None, None,
                                          vertical='General Transport')
    assert billable == set(ff_charges.groups_for('General Transport'))
    assert ff_charges.is_billable('line_haul', None, None,
                                  vertical='General Transport')


def test_forwarding_still_consults_the_incoterm():
    """The rule R1 established does not change because R5 widened it."""
    assert ff_charges.uses_incoterms('Freight Forwarding')
    seller = ff_charges.billable_groups('FOB', 'Shipper',
                                        vertical='Freight Forwarding')
    assert seller == {'origin', 'export_customs'}
    assert not ff_charges.is_billable(
        'freight', 'FOB', 'Shipper', vertical='Freight Forwarding')


def test_an_incoterm_is_still_demanded_where_it_applies():
    with pytest.raises(ff_charges.ChargeError):
        ff_charges.billable_groups('NONSENSE', 'Shipper',
                                   vertical='Freight Forwarding')


# ── one copy of everything ───────────────────────────────────────────
@pytest.mark.parametrize('absent', [
    'app/services/transport_rates.py',
    'app/services/transport_charges.py',
    'app/services/transport_handover.py',
    'app/services/quote_rates.py',
])
def test_there_is_no_second_copy_of_the_machinery(absent):
    assert not os.path.exists(os.path.join(_ROOT, absent)), \
        f'{absent} is a second copy of something ff_* already does'


def test_the_two_packs_no_longer_share_a_writer():
    """They did, while both were CRM-shaped. The real formats are
    not alike: Part A is consumed by a reader, one field name per
    row, and Part B is typed into three screens by a person. One
    writer would have to be wrong for one of its two readers."""
    src = open(os.path.join(_ROOT, 'app', 'services',
                            'job_handover.py')).read()
    assert 'ff_handover._write_xlsx' not in src
    assert 'Type into the TMS' in src


def test_the_job_payload_invents_nothing_now():
    """The contract arrived: docs/tms_import_formats.md, Part B.
    What replaces INVENTED is CANNOT_SUPPLY — real field names the
    CRM has no column for."""
    assert not hasattr(job_handover, 'INVENTED')
    assert 'procam_gst_reg_id' in job_handover.CANNOT_SUPPLY
    assert job_handover.THEIRS_TO_NUMBER == (
        'code', 'pcn_job_no', 'po_number')


@pytest.mark.parametrize('vertical,job_type', [
    ('General Transport', 'Road'),
    ('Heavy Transport', 'Road'),
    ('Project Freight', 'Sea'),
    ('Installation', 'Installation'),
    ('Warehousing', 'Warehouse'),
    ('Something Else', 'Road'),
])
def test_a_job_type_per_vertical(vertical, job_type):
    """ROAD, ODC, INSTALL and GENERAL were mine. These are the
    values the TMS's own worked examples use."""
    assert job_handover.job_type_for(vertical) == job_type


# ── the terms library ────────────────────────────────────────────────
@pytest.fixture
def terms():
    with flask_app.app_context():
        from app.models.quote_terms import QuoteTermSet
        QuoteTermSet.query.filter_by(vertical='R5 Test').delete()
        db.session.commit()
        yield
        QuoteTermSet.query.filter_by(vertical='R5 Test').delete()
        db.session.commit()


def test_publishing_supersedes_rather_than_edits(terms):
    with flask_app.app_context():
        one = quote_terms.publish('R5 Test', 'First wording.',
                                  actor='T')
        db.session.commit()
        two = quote_terms.publish('R5 Test', 'Second wording.',
                                  actor='T')
        db.session.commit()

        assert one.version == 1 and two.version == 2
        assert two.is_current and not one.is_current
        assert one.superseded_by_id == two.id
        assert one.body == 'First wording.', \
            'the old wording is what a client agreed to'
        assert quote_terms.current_for('R5 Test').id == two.id
        assert len(quote_terms.history_for('R5 Test')) == 2


def test_empty_terms_are_refused(terms):
    with flask_app.app_context():
        with pytest.raises(quote_terms.TermsError):
            quote_terms.publish('R5 Test', '   ', actor='T')
        with pytest.raises(quote_terms.TermsError):
            quote_terms.publish('', 'Something', actor='T')


# ── quote builder v2 ─────────────────────────────────────────────────
def _purge():
    """Everything this module makes, lines before quotes.

    Deleting a quote and leaving its lines behind is not a tidy-up
    with a loose end: SQLite reuses the freed row id, the next test
    module's quote is handed it, and that quote silently inherits
    these lines. It cost three failures in test_ff_rates — a margin
    of 220,000 where the test had built 130,000 — and the module
    passed on its own the whole time.
    """
    from app.models.quote import Quote, QuoteLine
    from app.models.quote_terms import (QuoteSpecialCondition,
                                        QuoteTermSet)

    ids = [q.id for q in Quote.query.filter(
        Quote.subject == 'r5').all()]
    if ids:
        QuoteLine.query.filter(QuoteLine.quote_id.in_(ids)).delete(
            synchronize_session=False)
        QuoteSpecialCondition.query.filter(
            QuoteSpecialCondition.quote_id.in_(ids)).delete(
                synchronize_session=False)
        Quote.query.filter(Quote.id.in_(ids)).delete(
            synchronize_session=False)
    QuoteTermSet.query.filter(
        QuoteTermSet.vertical.in_(('R5 Test', 'Heavy Transport'))
    ).delete(synchronize_session=False)
    for l in Lead.query.filter(Lead.source == 'r5').all():
        db.session.delete(l)
    db.session.commit()


@pytest.fixture
def quoted():
    from app.models.quote import Quote
    with flask_app.app_context():
        _purge()

        lead = Lead(source='r5', company='Haulage Co',
                    stage='Quoted', procam_vertical='Heavy Transport')
        db.session.add(lead)
        db.session.flush()
        q = Quote(quote_number=f'R5-{lead.id}', subject='r5',
                  lead_id=lead.id, currency='INR',
                  quote_date=date.today())
        db.session.add(q)
        db.session.commit()
        yield {'lead': lead.id, 'quote': q.id}
        _purge()


def test_a_line_must_use_a_group_this_vertical_has(quoted):
    from app.models.quote import Quote
    with flask_app.app_context():
        q = db.session.get(Quote, quoted['quote'])
        with pytest.raises(quote_build.QuoteBuildError) as exc:
            quote_build.add_line(q, service='Customs',
                                 charge_group='export_customs',
                                 amount=1000)
        assert 'Heavy Transport' in str(exc.value)
        line = quote_build.add_line(q, service='Trailer',
                                    charge_group='line_haul',
                                    amount=50000, cost_amount=40000)
        assert line.charge_group == 'line_haul'
        db.session.rollback()


def test_the_lead_mirrors_the_latest_quote(quoted):
    from app.models.quote import Quote
    with flask_app.app_context():
        q = db.session.get(Quote, quoted['quote'])
        quote_build.add_line(q, service='Trailer',
                             charge_group='line_haul', amount=50000)
        quote_build.add_line(q, service='Escort',
                             charge_group='escort_permits',
                             amount=5000)
        quote_build.recalculate(q, actor='T')
        db.session.commit()

        lead = db.session.get(Lead, quoted['lead'])
        assert float(q.total_amount) == 55000
        assert lead.quote_no == q.quote_number
        assert float(lead.quote_value_num) == 55000
        assert float(lead.quoted_amount_inr) == 55000
        assert lead.value_confirmed is True, \
            'the owner built the lines it adds up from'
        assert lead.value_source == 'quote'


def test_an_older_revision_does_not_pull_the_lead_backwards(quoted):
    from app.models.quote import Quote
    with flask_app.app_context():
        old = db.session.get(Quote, quoted['quote'])
        quote_build.add_line(old, service='Trailer',
                             charge_group='line_haul', amount=50000)
        quote_build.recalculate(old, actor='T')
        db.session.commit()

        newer = Quote(quote_number=f'R5-{quoted["lead"]}-B',
                      subject='r5', lead_id=quoted['lead'],
                      currency='INR', quote_date=date.today())
        db.session.add(newer)
        db.session.flush()
        quote_build.add_line(newer, service='Trailer',
                             charge_group='line_haul', amount=90000)
        quote_build.recalculate(newer, actor='T')
        db.session.commit()

        lead = db.session.get(Lead, quoted['lead'])
        assert float(lead.quote_value_num) == 90000

        # Editing the superseded one must not move the lead.
        quote_build.add_line(old, service='Extra',
                             charge_group='waiting', amount=1000)
        quote_build.recalculate(old, actor='T')
        db.session.commit()
        lead = db.session.get(Lead, quoted['lead'])
        assert float(lead.quote_value_num) == 90000, \
            'an old revision rewrote the lead'


def test_the_terms_version_is_stamped_once(quoted, terms):
    from app.models.quote import Quote
    with flask_app.app_context():
        quote_terms.publish('Heavy Transport', 'Version one.',
                            actor='T')
        db.session.commit()
        q = db.session.get(Quote, quoted['quote'])
        quote_build.recalculate(q, actor='T')
        db.session.commit()
        first = q.terms_version_id
        assert first

        quote_terms.publish('Heavy Transport', 'Version two.',
                            actor='T')
        db.session.commit()
        quote_build.recalculate(q, actor='T')
        db.session.commit()
        assert q.terms_version_id == first, \
            'the client has the old wording in their hand'
        from app.models.quote_terms import QuoteTermSet
        QuoteTermSet.query.filter_by(
            vertical='Heavy Transport').delete()
        db.session.commit()


def test_special_conditions_are_numbered_and_refuse_to_be_empty(
        quoted):
    from app.models.quote import Quote
    with flask_app.app_context():
        q = db.session.get(Quote, quoted['quote'])
        a = quote_terms.add_condition(q, 'Crane by client.', actor='T')
        b = quote_terms.add_condition(q, 'Night work excluded.',
                                      actor='T')
        db.session.commit()
        assert (a.seq, b.seq) == (1, 2)
        with pytest.raises(quote_terms.TermsError):
            quote_terms.add_condition(q, '   ', actor='T')
        printed = quote_terms.for_print(q)
        assert [c['text'] for c in printed['special']] == [
            'Crane by client.', 'Night work excluded.']


def test_a_quote_with_no_published_terms_says_so(quoted, terms):
    from app.models.quote import Quote
    with flask_app.app_context():
        from app.models.quote_terms import QuoteTermSet
        QuoteTermSet.query.filter_by(
            vertical='Heavy Transport').delete()
        db.session.commit()
        q = db.session.get(Quote, quoted['quote'])
        printed = quote_terms.for_print(q)
        assert printed['standard']['missing'] is True


# ── sending ──────────────────────────────────────────────────────────
def test_the_send_path_goes_through_the_outbox():
    """Not a second SMTP connection: the caps, retries, dedupe and
    kill switch all live in the outbox."""
    src = open(os.path.join(_ROOT, 'app', 'services',
                            'quote_send.py')).read()
    assert 'outbox.enqueue' in src
    assert 'smtplib' not in src
    assert 'SMTP(' not in src


def test_a_full_day_refuses_before_queueing(monkeypatch):
    """Told now, rather than finding out tomorrow that it queued."""
    from app.services import outbox
    monkeypatch.setattr(outbox, 'day_room', lambda *a, **k: 0)
    with flask_app.app_context():
        with pytest.raises(quote_send.SendRefused) as exc:
            quote_send.send_quote(object(), to='x@example.com')
        assert 'daily recipient cap' in str(exc.value)


def test_a_vendor_with_no_address_is_refused_by_name():
    class _R:
        id = 1
        vendor_email = None
        vendor_name = 'Blue Water Shipping'
    with flask_app.app_context():
        with pytest.raises(quote_send.SendRefused) as exc:
            quote_send.send_rate_request(_R())
        assert 'Blue Water Shipping' in str(exc.value)


def test_forwarding_masks_by_default_and_transport_does_not():
    class _L:
        procam_vertical = 'Freight Forwarding'

    class _T:
        procam_vertical = 'Heavy Transport'
    assert quote_send._masks_by_default(_L(), None) is True
    assert quote_send._masks_by_default(_T(), None) is False


def test_the_pdf_escapes_the_terms_it_prints():
    """reportlab reads markup: "Rates & charges" would raise inside
    doc.build, on the one page a client sees."""
    from app.quote.routes import _esc
    assert _esc('Rates & charges < 24h') == \
        'Rates &amp; charges &lt; 24h'


# ── reading terms off the quotes we have actually sent ───────────────
def _terms_mod():
    import importlib.util as u
    spec = u.spec_from_file_location(
        'terms_from_quotes',
        os.path.join(_ROOT, 'scripts', 'crm2_terms_from_quotes.py'))
    mod = u.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.mark.parametrize('text,why', [
    ('The LSP shall deploy vehicles with valid RC. Our offer is '
     'attached.', 'the client binding a supplier'),
    ('You are to quote clearly the number of trailers. Procam '
     'Logistics Pvt Ltd.', 'the client instructing a bidder'),
    ('Please submit your offer in a sealed envelope. Procam '
     'Logistics Private Limited', 'a tender'),
    ('Bidders must comply. We are pleased to quote.',
     'a tender that also says our words'),
])
def test_a_clients_tender_is_not_our_quote(text, why):
    """The first run scored on "mentions Procam" and built a
    proposal out of TMIL's tender and L&T's LSP contract. Those
    name Procam on every page."""
    mod = _terms_mod()
    score, hits, _why_not = mod._score(text)
    assert score == 0, f'{why}: scored {score} on {hits}'


def test_our_own_offer_is_recognised():
    mod = _terms_mod()
    score, hits, _ = mod._score(
        'We are pleased to quote as below. Our offer is valid for 15 '
        'days. Payment terms: 30 days. Procam Logistics Pvt Ltd.')
    assert score >= 4 and hits


def test_a_veto_outranks_any_number_of_positives():
    """A tender can carry our letterhead on a covering page."""
    mod = _terms_mod()
    score, _h, why = mod._score(
        'Procam Logistics Pvt Ltd. Our quotation. We are pleased to '
        'quote. Our charges. The LSP shall provide tracking.')
    assert score == 0 and 'LSP' in why


@pytest.mark.parametrize('line,keep', [
    ('The present offer is valid for 15 days from the date hereof',
     True),
    ('All risk insurance to cover the full value of the shipment',
     True),
    ('Axles lines (upto 10) with', False),
    ('4921 / 4923 and equivalent 1500 3000 4500', False),
    ('Procam Logistics Pvt. Ltd.,', False),
    ('You should provide a proper loading plan with all packages',
     False),
    ('The Contractor shall not transship any goods without consent',
     False),
    ('Transporter shall provide a tracking report every day',
     False),
])
def test_a_clause_is_a_sentence_of_terms_not_a_table_row(line, keep):
    assert _terms_mod()._is_clause(line) is keep


def test_the_proposal_script_publishes_nothing():
    """It reads. `quote_terms.publish()` is a separate, deliberate
    act, and a proposal that quietly became a publication would be
    the worst version of this task.

    Checked against the parsed code rather than the text: the module
    docstring says the word "publish" twice, on purpose, and a
    grep-shaped test fails on its own explanation.
    """
    import ast

    tree = ast.parse(open(os.path.join(
        _ROOT, 'scripts', 'crm2_terms_from_quotes.py')).read())
    called = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            called.add(ast.unparse(node.func))
    for banned in ('quote_terms.publish', 'db.session.commit',
                   'db.session.add', 'db.session.delete'):
        assert banned not in called, \
            f'a proposal script must not call {banned}()'
    assert 'db.session.execute' in called, \
        'it does have to read the attachments'


def test_a_word_terms_annexure_is_read_tables_and_all():
    """A terms annexure is very often a two-column table, and
    reading only document.paragraphs returns the headings and none
    of the clauses."""
    import docx
    from app.services import attachment_text

    assert attachment_text.word_supported()
    path = os.path.join(tempfile.mkdtemp(), 'terms.docx')
    d = docx.Document()
    d.add_paragraph('Terms and Conditions')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = '1'
    t.cell(0, 1).text = 'The present offer is valid for 15 days.'
    t.cell(1, 0).text = '2'
    t.cell(1, 1).text = 'All risk insurance is to the client.'
    d.save(path)

    text = attachment_text.extract(path)
    assert 'Terms and Conditions' in text
    assert 'valid for 15 days' in text, 'the table was not read'
    assert 'All risk insurance' in text


def test_a_merged_row_is_not_repeated_once_per_cell():
    import docx
    from app.services import attachment_text

    path = os.path.join(tempfile.mkdtemp(), 'merged.docx')
    d = docx.Document()
    t = d.add_table(rows=1, cols=3)
    for c in t.rows[0].cells:
        c.text = 'Payment within 30 days'
    d.save(path)
    assert attachment_text.extract(path).count(
        'Payment within 30 days') == 1


def test_docx_is_in_the_readable_list():
    from app.services import attachment_text
    assert '.docx' in attachment_text.READABLE
    assert '.doc' not in attachment_text.READABLE, \
        'the legacy binary format needs a reader this server has not'
