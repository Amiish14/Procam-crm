"""
CRM-VERTICAL-FIX-01 — saving a lead must not move its vertical.

The lead edit form offered four verticals. Leads carry eleven. A
`<select>` whose options do not contain the record's value selects
option zero, and the page posted option zero as though somebody had
chosen it. 443 of 4,505 leads are suspected, and the audit trail
records the rewrites by name and date:

    12204  Renew Power   Freight Forwarding -> Heavy Transport  15 Oct
    12148  Tmilltd       Customs            -> Heavy Transport  12 Oct
    12218  Prysmian      Project Freight    -> Heavy Transport  16 Oct

Two defences, because either alone can be defeated:

    the form   renders the record's own value, always, whether or not
               the master list knows about it;
    the server accepts a change only when the posted value differs
               from what the form was rendered with — the difference
               between a choice and a fallback.
"""
import json
import os
import sys
import tempfile
from datetime import datetime, timedelta

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

os.environ['URL_PREFIX'] = ''
os.environ['SESSION_COOKIE_SECURE'] = 'false'
os.environ.setdefault('SECRET_KEY', 'vertfix-test-secret')
os.environ.setdefault('ADMIN_INITIAL_PASSWORD', 'VertFixTest12345')
os.environ.setdefault(
    'DATABASE_URL',
    'sqlite:///' + os.path.join(tempfile.mkdtemp(), 'vertfix.db'))

import importlib                                          # noqa: E402
_main = importlib.import_module('app')
importlib.import_module('app.models')
flask_app, db = _main.app, _main.db
Employee, Lead, Company = _main.Employee, _main.Lead, _main.Company
flask_app.config['WTF_CSRF_ENABLED'] = False

from app.models.access import AccessProfile, DataScope    # noqa: E402

ADMIN = 'VFADMIN'
TAG = 'VF '

#: Every vertical a lead actually carries, not the four the form used
#: to offer. The four were Heavy Transport, Project Freight,
#: Warehousing and Installation; the other seven are the ones that
#: were being destroyed.
ALL_VERTICALS = (
    'Heavy Transport', 'Project Freight', 'Warehousing', 'Installation',
    'Freight Forwarding', 'Customs', 'Sea Freight', 'Air Freight',
    'General Transport', 'Heavy Cargo', 'Other',
)


def _emp(code, name, *, role='admin', scope=DataScope.ALL, perms=()):
    e = Employee.query.filter_by(emp_code=code).first() or Employee(emp_code=code)
    e.name, e.is_active, e.must_change_pw = name, True, False
    e.email = f'{code.lower()}@procamgroup.in'
    e.role, e.session_version = role, 0
    e.is_super_admin = (role == 'admin')
    e.vertical = 'All'
    db.session.add(e)
    db.session.flush()
    AccessProfile.query.filter_by(emp_code=code).delete()
    db.session.add(AccessProfile(emp_code=code, data_scope=scope,
                                 perms=list(perms)))
    return e


def _wipe():
    from app.models.audit import AuditEvent

    ids = [l.id for l in Lead.query.filter(Lead.company.like(TAG + '%'))]
    if ids:
        # The audit rows go too. SQLite reuses rowids after a delete,
        # so a lead created by the next test inherits the previous
        # one's history and the restore reads a rewrite that belongs
        # to a lead that no longer exists.
        AuditEvent.query.filter(
            AuditEvent.entity_type.in_(('Lead', 'lead')),
            AuditEvent.entity_id.in_([str(i) for i in ids])).delete(
                synchronize_session=False)
        Lead.query.filter(Lead.id.in_(ids)).delete(synchronize_session=False)
    db.session.commit()


@pytest.fixture()
def world():
    with flask_app.app_context():
        db.create_all()
        _wipe()
        _emp(ADMIN, 'Vertical Fix Admin')
        made = {}
        for vertical in ALL_VERTICALS:
            lead = Lead(company=TAG + vertical, assigned_to=ADMIN,
                        assigned_name=ADMIN, stage='Quoted',
                        procam_vertical=vertical,
                        followup_date=None,
                        created_at=datetime.utcnow() - timedelta(days=5))
            db.session.add(lead)
            db.session.flush()
            made[vertical] = lead.id
        db.session.commit()
        yield made
        _wipe()


def _client(code=ADMIN, role='admin'):
    c = flask_app.test_client()
    with c.session_transaction() as s:
        s.update(emp_code=code, name=code, role=role, vertical='All')
    return c


def _save(lid, body):
    return _client().put(f'/api/leads/{lid}', data=json.dumps(body),
                         content_type='application/json')


def _vertical(lid):
    with flask_app.app_context():
        return db.session.get(Lead, lid).procam_vertical


# ── 1. the form fix is live and not behind a flag ────────────────────
def test_the_form_builds_its_list_from_the_master_not_a_literal():
    page = open(os.path.join(_ROOT, 'templates', 'app.html'),
                encoding='utf-8').read()
    assert 'verticalOptions(l.procam_vertical)' in page
    # Scoped to the lead form. The same four-item literal also backs
    # the dashboard's vertical doughnut, which is a display limit —
    # that chart shows four of the eleven verticals — and not a data
    # rewrite. Widening this assertion to the whole file would make
    # it fail for the wrong reason.
    start = page.index('Procam vertical</label>')
    form = page[start:start + 400]
    assert "'Heavy Transport'" not in form, 'the hard-coded list is back'
    assert 'verticalOptions(' in form


def test_the_fix_is_not_behind_a_feature_flag():
    """It must be on for everyone. A silent data rewrite is not
    something to roll out gradually."""
    page = open(os.path.join(_ROOT, 'templates', 'app.html'),
                encoding='utf-8').read()
    start = page.index('function verticalOptions(')
    end = page.index('}', page.index('return list.map', start))
    body = page[start:end]
    for marker in ('feature(', 'FEATURE_', 'flags.on'):
        assert marker not in body, f'{marker} guards the vertical fix'


def test_each_select_sends_what_it_was_rendered_with():
    page = open(os.path.join(_ROOT, 'templates', 'app.html'),
                encoding='utf-8').read()
    for sel in ('lVt', 'lAsgn', 'lAsgn2'):
        assert f'id="{sel}" data-was=' in page, sel
    for key in ('procam_vertical_was', 'assigned_to_was',
                'secondary_owner_was'):
        assert key in page, key


# ── 2. editing something else must not move the vertical ─────────────
@pytest.mark.parametrize('vertical', ALL_VERTICALS)
def test_editing_an_unrelated_field_leaves_the_vertical_alone(world,
                                                              vertical):
    """The exact shape of the bug: the user changes a follow-up date,
    and the select posts Heavy Transport because the lead's real
    vertical was not among its options."""
    lid = world[vertical]
    res = _save(lid, {
        'followup': '2026-12-01',
        'next_action': 'Call the buyer',
        # what a defective or cached form posts: option zero, with
        # the rendered value saying it was never touched
        'procam_vertical': 'Heavy Transport',
        'procam_vertical_was': 'Heavy Transport',
    })
    assert res.status_code == 200
    assert _vertical(lid) == vertical, (
        f'{vertical} was rewritten to {_vertical(lid)}')
    with flask_app.app_context():
        assert str(db.session.get(Lead, lid).followup_date) == '2026-12-01', \
            'the edit the user actually made must still be saved'


@pytest.mark.parametrize('vertical', ALL_VERTICALS)
def test_the_corrected_form_posts_the_real_value_and_nothing_moves(world,
                                                                   vertical):
    """What the fixed page sends: the lead's own value, untouched."""
    lid = world[vertical]
    res = _save(lid, {'next_action': 'Chase',
                      'procam_vertical': vertical,
                      'procam_vertical_was': vertical})
    assert res.status_code == 200
    assert _vertical(lid) == vertical


def test_a_vertical_the_user_really_changed_is_saved(world):
    lid = world['Freight Forwarding']
    res = _save(lid, {'procam_vertical': 'Customs',
                      'procam_vertical_was': 'Freight Forwarding'})
    assert res.status_code == 200
    assert _vertical(lid) == 'Customs'


def test_a_page_opened_before_someone_else_changed_it_does_not_clobber(world):
    lid = world['Customs']
    with flask_app.app_context():
        db.session.get(Lead, lid).procam_vertical = 'Sea Freight'
        db.session.commit()
    res = _save(lid, {'procam_vertical': 'Heavy Transport',
                      'procam_vertical_was': 'Customs'})
    assert res.status_code == 200
    assert _vertical(lid) == 'Sea Freight'
    assert any('somebody else changed it' in w
               for w in res.get_json().get('warnings', []))


def test_a_cached_page_with_no_rendered_value_is_refused(world):
    """A page served before this guard cannot tell a choice from a
    default, so the CRM does not guess."""
    lid = world['Project Freight']
    res = _save(lid, {'procam_vertical': 'Heavy Transport'})
    assert res.status_code == 200
    assert _vertical(lid) == 'Project Freight'
    assert any('loaded before the form was corrected' in w
               for w in res.get_json().get('warnings', []))


def test_the_rest_of_the_save_survives_a_refused_vertical(world):
    lid = world['Heavy Cargo']
    res = _save(lid, {'procam_vertical': 'Heavy Transport',
                      'next_action': 'Send the revised offer',
                      'followup': '2026-11-20'})
    assert res.status_code == 200
    with flask_app.app_context():
        lead = db.session.get(Lead, lid)
        assert lead.procam_vertical == 'Heavy Cargo'
        assert lead.next_action == 'Send the revised offer'
        assert str(lead.followup_date) == '2026-11-20'


# ── 3. the same shape on the other selects ───────────────────────────
def test_the_owner_select_carries_an_owner_who_has_left():
    """The same defect, fixed in the form rather than the API.

    `lAsgn` has no blank first option, and for anybody without
    admin.employees `/api/employees` returns active staff only — 29
    of 131. A lead owned by somebody who has left matched no option,
    so the select fell back to option zero and saving reassigned the
    lead to whoever was first.

    Guarded in the template and not on the server because
    `assigned_to` has other callers — the triage screen, the
    integration API, the assignment scripts — none of which send a
    rendered value, and all of which would break.
    """
    page = open(os.path.join(_ROOT, 'templates', 'app.html'),
                encoding='utf-8').read()
    assert 'const _ownOpt=' in page
    start = page.index('const _ownOpt=')
    block = page[start:page.index('const sOpts=', start) + 400]
    assert 'no longer active' in block
    assert '_ownOpt(l.assigned_to' in block
    assert '_ownOpt(l.secondary_owner' in block


def test_the_owner_fields_are_not_guarded_on_the_server():
    """Deliberately. Narrowing this guard to the one field with the
    proven defect is what keeps the other callers working."""
    assert set(_main._TOUCH_GUARDED) == {'procam_vertical'}


def test_a_real_reassignment_still_works(world):
    lid = world['Other']
    with flask_app.app_context():
        _emp('VFNEW', 'New Owner', role='user')
        db.session.commit()
    res = _save(lid, {'assigned_to': 'VFNEW',
                      'reassignment_reason': 'Different geography'})
    assert res.status_code == 200
    with flask_app.app_context():
        assert db.session.get(Lead, lid).assigned_to == 'VFNEW'


def test_the_stage_widget_is_not_a_select_and_cannot_default(world):
    """`ldPipe` reads its value from the lead, so there is no option
    zero to fall back to."""
    page = open(os.path.join(_ROOT, 'templates', 'app.html'),
                encoding='utf-8').read()
    assert 'id="ldPipe"' in page
    assert 'g(\'ldPipe\')?.dataset.stage' in page
    assert '<select class="fmc" id="lStg"' not in page


def test_the_account_form_cannot_write_a_vertical_at_all(world):
    """api_update_company's field list has no vertical in it, so the
    account screen cannot reproduce this bug."""
    import inspect
    src = inspect.getsource(_main.api_update_company)
    body = src[src.index("for f in ("):src.index("db.session.commit()")]
    assert 'vertical' not in body


# ── 4. the repair workbook ───────────────────────────────────────────
def _audit(lead_id, frm, to, *, actor, when, clear=True):
    from app.models.audit import AuditEvent

    # Clear anything already filed against this id first, on the
    # first row only — a fixture building a history of two changes
    # must not wipe its own earlier one. The suite
    # shares one database, other modules delete their leads without
    # deleting their audit rows, and SQLite reuses rowids — so a lead
    # created here inherits a stranger's history and the restore
    # reads a change that belongs to a record that no longer exists.
    if clear:
        AuditEvent.query.filter(
            AuditEvent.entity_type.in_(('Lead', 'lead')),
            AuditEvent.entity_id == str(lead_id)).delete(
                synchronize_session=False)

    db.session.add(AuditEvent(
        occurred_at=when, actor=actor, action='lead.update',
        entity_type='Lead', entity_id=str(lead_id),
        old_value={'procam_vertical': frm},
        new_value={'procam_vertical': to}))


@pytest.fixture()
def damaged(world):
    """Three leads, one per population."""
    from app.services import vertical_audit as va

    base = datetime(2026, 10, 1, 9, 0)
    with flask_app.app_context():
        made = {}
        for key, project in (
                ('recorded', 'Customs clearance, bill of entry at Nhava Sheva'),
                ('review', 'FCL shipment, port of loading Shanghai'),
                ('changed', 'Customs clearance and CHA work')):
            lead = Lead(company=TAG + key, assigned_to=ADMIN,
                        assigned_name=ADMIN, stage='Quoted',
                        procam_vertical=va.TARGET, project=project,
                        created_at=base, updated_at=datetime.utcnow())
            db.session.add(lead)
            db.session.flush()
            made[key] = lead.id

        _audit(made['recorded'], 'Freight Forwarding', va.TARGET,
               actor='EMP3592024', when=base)
        # rewritten, then set again by somebody else afterwards
        _audit(made['changed'], 'Customs', va.TARGET,
               actor='EMP3592024', when=base)
        _audit(made['changed'], va.TARGET, 'Project Freight',
               actor='DIR12010', when=base + timedelta(days=3),
               clear=False)
        db.session.commit()
        yield made


def test_the_three_populations_are_separated(damaged):
    from app.services import vertical_audit as va
    with flask_app.app_context():
        data = va.survey(since='2026-09-01')
    ids = {k: [r['lead_id'] for r in data[k]]
           for k in ('recorded', 'review', 'changed_since')}
    assert damaged['recorded'] in ids['recorded']
    assert damaged['review'] in ids['review']
    assert damaged['changed'] in ids['changed_since']
    # and never in more than one
    assert damaged['changed'] not in ids['recorded']
    assert damaged['recorded'] not in ids['review']


def test_a_recorded_restore_is_the_audited_old_value_not_a_guess(damaged):
    from app.services import vertical_audit as va
    with flask_app.app_context():
        data = va.survey(since='2026-09-01')
    row = next(r for r in data['recorded']
               if r['lead_id'] == damaged['recorded'])
    assert row['restore_to'] == 'Freight Forwarding'
    assert row['rewritten_by'] == 'EMP3592024'
    # the keywords say Customs; the audit trail says Freight
    # Forwarding, and the audit trail wins. The disagreement is
    # reported rather than resolved.
    assert row['looks_like'] == 'Customs'
    assert row['agrees'] == 'differs'


def test_a_lead_changed_after_the_rewrite_is_excluded(damaged):
    """Restoring it would undo a decision somebody took since."""
    from app.services import vertical_audit as va
    with flask_app.app_context():
        data = va.survey(since='2026-09-01')
    row = next(r for r in data['changed_since']
               if r['lead_id'] == damaged['changed'])
    assert row['was'] == 'Customs'
    assert row['changed_to'] == 'Project Freight'
    assert row['changed_by'] == 'DIR12010'


def test_review_rows_carry_the_evidence_and_no_answer(damaged):
    from app.services import vertical_audit as va
    with flask_app.app_context():
        data = va.survey(since='2026-09-01')
    row = next(r for r in data['review'] if r['lead_id'] == damaged['review'])
    assert row['looks_like'] == 'Sea Freight'
    assert row['evidence']
    assert 'restore_to' not in row, 'the CRM must not invent a restore'


def test_the_survey_writes_nothing(damaged):
    from app.services import vertical_audit as va
    with flask_app.app_context():
        before = [(l.id, l.procam_vertical) for l in Lead.query.all()]
        va.survey(since='2026-09-01')
        after = [(l.id, l.procam_vertical) for l in Lead.query.all()]
    assert before == after


def test_the_workbook_has_the_three_sheets(damaged, tmp_path):
    import subprocess
    out = tmp_path / 'repair.xlsx'
    with flask_app.app_context():
        url = str(db.engine.url)
    res = subprocess.run(
        [sys.executable, os.path.join(_ROOT, 'scripts',
                                      'vertical_repair_workbook.py'),
         '--out', str(out)],
        cwd=_ROOT, capture_output=True, text=True,
        env=dict(os.environ, URL_PREFIX='', DATABASE_URL=url))
    assert res.returncode == 0, res.stderr[-800:]
    assert out.exists()

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ['About', 'Recorded', 'Review', 'Changed since']
    assert wb['Review'][2][11].value == 'correct vertical'
    assert wb['Recorded'][2][6].value == 'restore_to'


# ── Stage B: putting them back ───────────────────────────────────────
def _run(script, *args):
    """Run a repair script against the database this process is on.

    The URL comes from the live engine, not from os.environ. Several
    test modules assign DATABASE_URL at import, so the variable ends
    up naming whichever module imported last while the engine stays
    bound to whichever imported first — and a subprocess trusting the
    variable reads a different database than the fixtures wrote to.
    """
    import subprocess

    with flask_app.app_context():
        url = str(db.engine.url)
    return subprocess.run(
        [sys.executable, os.path.join(_ROOT, 'scripts', script), *args],
        cwd=_ROOT, capture_output=True, text=True,
        env=dict(os.environ, URL_PREFIX='', DATABASE_URL=url))


@pytest.fixture()
def to_repair(world):
    """One lead of each kind the restore has to tell apart."""
    from app.models.email_admin import VerticalAlias
    from app.services import vertical_audit as va

    base = datetime(2026, 10, 1, 9, 0)
    with flask_app.app_context():
        for alias, real in (('Transportation', 'Heavy Transport'),
                            ('Project Logistics', 'Project Freight')):
            if not VerticalAlias.query.filter_by(alias=alias).first():
                db.session.add(VerticalAlias(alias=alias, vertical=real))
        made = {}
        for key, was in (('restore', 'Heavy Cargo'),
                         ('remap', 'Project Logistics'),
                         ('skip_alias', 'Transportation'),
                         ('skip_deliberate', 'Project Freight')):
            lead = Lead(company=TAG + 'B ' + key, assigned_to=ADMIN,
                        assigned_name=ADMIN, stage='Quoted',
                        procam_vertical=va.TARGET,
                        project='customs clearance at Nhava Sheva',
                        created_at=base, updated_at=datetime.utcnow())
            db.session.add(lead)
            db.session.flush()
            made[key] = lead.id
            _audit(lead.id, was, va.TARGET, actor='DIR12010', when=base)
        db.session.commit()
        yield made


def test_the_dry_run_writes_nothing(to_repair):
    from app.services import vertical_audit as va
    res = _run('vertical_restore.py', '--check')
    assert res.returncode == 0, res.stderr[-800:]
    assert 'DRY RUN' in res.stdout
    with flask_app.app_context():
        for lid in to_repair.values():
            assert db.session.get(Lead, lid).procam_vertical == va.TARGET


def test_apply_restores_remaps_and_skips_the_right_ones(to_repair):
    from app.services import vertical_audit as va
    res = _run('vertical_restore.py', '--apply')
    assert res.returncode == 0, res.stderr[-800:]
    with flask_app.app_context():
        got = {k: db.session.get(Lead, v).procam_vertical
               for k, v in to_repair.items()}
    assert got['restore'] == 'Heavy Cargo'
    # the dead label becomes what it means, not the label itself
    assert got['remap'] == 'Project Freight'
    # an alias of the value it already carries: untouched
    assert got['skip_alias'] == va.TARGET
    # the form offered Project Freight, so the defect cannot explain it
    assert got['skip_deliberate'] == va.TARGET


def test_the_four_the_form_offered_are_listed_with_their_exact_value(
        to_repair):
    res = _run('vertical_restore.py', '--check')
    assert 'the defect cannot explain it' in res.stdout
    assert "'Project Freight'" in res.stdout, \
        'the stored value must be shown exactly, quoting included'


def test_every_written_lead_gets_the_reason_on_its_audit_row(to_repair):
    from app.models.audit import AuditEvent
    import scripts.vertical_restore as vr  # noqa: F401  (for REASON)

    _run('vertical_restore.py', '--apply')
    with flask_app.app_context():
        rows = (AuditEvent.query
                .filter(AuditEvent.action == 'lead.vertical_restored')
                .all())
        touched = {int(r.entity_id) for r in rows}
        assert to_repair['restore'] in touched
        assert to_repair['remap'] in touched
        assert to_repair['skip_alias'] not in touched
        assert to_repair['skip_deliberate'] not in touched
        for r in rows:
            assert r.reason == ('restored after form defect, '
                                'CRM-VERTICAL-FIX-01')


def test_a_lead_someone_moved_since_is_left_alone(to_repair):
    with flask_app.app_context():
        db.session.get(Lead, to_repair['restore']).procam_vertical = \
            'Sea Freight'
        db.session.commit()
    res = _run('vertical_restore.py', '--apply')
    assert res.returncode == 0
    with flask_app.app_context():
        assert db.session.get(Lead, to_repair['restore']).procam_vertical \
            == 'Sea Freight'


def test_running_it_twice_changes_nothing_the_second_time(to_repair):
    _run('vertical_restore.py', '--apply')
    with flask_app.app_context():
        after_one = {k: db.session.get(Lead, v).procam_vertical
                     for k, v in to_repair.items()}
    _run('vertical_restore.py', '--apply')
    with flask_app.app_context():
        after_two = {k: db.session.get(Lead, v).procam_vertical
                     for k, v in to_repair.items()}
    assert after_one == after_two


def test_a_failure_part_way_through_rolls_the_whole_thing_back(to_repair):
    """One transaction means one transaction. A half-applied repair
    is worse than none: nobody could tell which half."""
    from app.services import vertical_audit as va
    from app.services import audit as audit_mod

    calls = {'n': 0}
    real = audit_mod.record

    def explode(*a, **kw):
        calls['n'] += 1
        if calls['n'] == 2:
            raise RuntimeError('simulated failure part way through')
        return real(*a, **kw)

    with flask_app.app_context():
        before = {k: db.session.get(Lead, v).procam_vertical
                  for k, v in to_repair.items()}

    import scripts.vertical_restore as vr
    original = vr.main
    audit_mod.record = explode
    try:
        with flask_app.app_context():
            import sys as _sys
            argv = _sys.argv
            _sys.argv = ['vertical_restore.py', '--apply']
            try:
                vr.main()
            except SystemExit as exc:
                assert exc.code == 1
            finally:
                _sys.argv = argv
    finally:
        audit_mod.record = real
        vr.main = original

    with flask_app.app_context():
        after = {k: db.session.get(Lead, v).procam_vertical
                 for k, v in to_repair.items()}
    assert after == before, 'a failed run must leave nothing behind'


# ── the split review files ───────────────────────────────────────────
def test_the_review_splits_into_one_file_per_suggestion(damaged, tmp_path):
    res = _run('vertical_review_split.py', '--out-dir', str(tmp_path))
    assert res.returncode == 0, res.stderr[-800:]
    made = sorted(p.name for p in tmp_path.glob('vertical_review_*.xlsx'))
    assert made, res.stdout[-600:]
    assert any('sea_freight' in n for n in made)


def test_each_file_carries_a_dropdown_of_valid_verticals(damaged, tmp_path):
    from openpyxl import load_workbook

    _run('vertical_review_split.py', '--out-dir', str(tmp_path))
    path = next(tmp_path.glob('vertical_review_*.xlsx'))
    wb = load_workbook(path)
    ws = wb['Review']

    assert [c.value for c in ws[2]] == [
        'lead', 'client', 'lead title', 'owner', 'confidence',
        'evidence', 'correct vertical']

    assert 'Verticals' in wb.sheetnames
    offered = [c.value for c in wb['Verticals']['A'][1:] if c.value]
    assert 'Sea Freight' in offered and 'Customs' in offered
    assert 'Management' not in offered, 'not a service a lead can be in'
    assert 'All' not in offered
    # Dead labels are not offered: picking Transportation believing
    # it differs from Heavy Transport is the confusion the alias
    # table exists to end. The applier still accepts them.
    assert 'Transportation' not in offered
    assert 'Project Logistics' not in offered

    dvs = list(ws.data_validations.dataValidation)
    assert dvs, 'no dropdown on the answer column'
    assert dvs[0].type == 'list'
    assert 'Verticals' in dvs[0].formula1

    # allowBlank is read from the file, not from openpyxl's loader:
    # the loader drops the flag, so asserting on the object tests a
    # reader quirk while Excel reads the XML and honours it.
    import zipfile
    with zipfile.ZipFile(path) as z:
        sheet = next(n for n in z.namelist()
                     if n.startswith('xl/worksheets/sheet'))
        xml = z.read(sheet).decode()
    i = xml.find('<dataValidation ')
    assert i != -1, 'no dataValidation element in the file'
    assert 'allowBlank="1"' in xml[i:i + 400], \
        'blank must stay a valid answer'


def test_high_confidence_rows_come_first(damaged, tmp_path):
    from openpyxl import load_workbook

    _run('vertical_review_split.py', '--out-dir', str(tmp_path))
    for path in tmp_path.glob('vertical_review_*.xlsx'):
        ws = load_workbook(path)['Review']
        hdr = [c.value for c in ws[2]]
        col = hdr.index('confidence')
        seen = [str(r[col].value or '').lower()
                for r in ws.iter_rows(min_row=3) if r[0].value]
        ranks = [{'high': 0, 'medium': 1, 'low': 2}.get(c, 3) for c in seen]
        assert ranks == sorted(ranks), f'{path.name}: {seen}'


def test_the_dropdown_offers_exactly_what_the_applier_accepts(damaged):
    """A sheet that suggests a value the applier then refuses is the
    worst of both."""
    from openpyxl import load_workbook  # noqa: F401
    import importlib

    with flask_app.app_context():
        from app.services import vertical_audit as va
        offered = set(va.acceptable_verticals(for_choice=True))
    mod = importlib.import_module('scripts.vertical_apply_review')
    with flask_app.app_context():
        accepted = mod.known_verticals()
    assert offered <= accepted


def test_a_filled_in_split_file_applies(damaged, tmp_path):
    """The whole round trip: split, fill one row in, apply it."""
    from openpyxl import load_workbook
    from app.services import vertical_audit as va

    _run('vertical_review_split.py', '--out-dir', str(tmp_path))
    path = next(p for p in tmp_path.glob('vertical_review_*.xlsx'))
    wb = load_workbook(path)
    ws = wb['Review']
    hdr = [c.value for c in ws[2]]
    lead_id = ws.cell(row=3, column=hdr.index('lead') + 1).value
    ws.cell(row=3, column=hdr.index('correct vertical') + 1,
            value='Sea Freight')
    # The applier reads a sheet called Review with its own header
    # shape, so the split file has to be readable by it as it stands.
    wb.save(path)

    res = _run('vertical_apply_review.py', '--file', str(path),
               '--by', 'TESTER', '--apply')
    assert res.returncode == 0, res.stderr[-800:]
    with flask_app.app_context():
        assert db.session.get(Lead, lead_id).procam_vertical == 'Sea Freight'


# ── CRM-VERTICAL-MASTER-01 Stage A ───────────────────────────────────
@pytest.fixture()
def master_world(world):
    """Values of every kind the survey has to tell apart."""
    from app.models.email_admin import VerticalAlias, VerticalHead

    with flask_app.app_context():
        for alias, real in (('Transportation', 'Heavy Transport'),
                            ('Project Logistics', 'Project Freight')):
            if not VerticalAlias.query.filter_by(alias=alias).first():
                db.session.add(VerticalAlias(alias=alias, vertical=real))
        for v in ('Heavy Transport', 'Project Freight', 'Management'):
            if not VerticalHead.query.filter_by(vertical=v).first():
                db.session.add(VerticalHead(vertical=v, active=True))
        db.session.flush()
        # a head who has left, and a vertical with nobody
        gone = _emp('VFLEFT', 'Departed Head', role='user')
        gone.is_active = False
        hv = VerticalHead.query.filter_by(vertical='Heavy Transport').first()
        hv.head_emp_code = 'VFLEFT'
        VerticalHead.query.filter_by(vertical='Management').first() \
            .head_emp_code = None

        for value in ('Heavy Transport', 'Transportation', 'Heavy Cargo',
                      'heavy  cargo', None):
            for _i in range(2):
                db.session.add(Lead(
                    company=TAG + 'M ' + (value or 'blank'),
                    assigned_to=ADMIN, stage='Quoted',
                    procam_vertical=value,
                    created_at=datetime.utcnow() - timedelta(days=3),
                    updated_at=datetime.utcnow()))
        db.session.commit()
        yield


def test_the_survey_separates_certain_from_needs_a_decision(master_world,
                                                            tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / 'master.xlsx'
    res = _run('vertical_master_survey.py', '--out', str(out))
    assert res.returncode == 0, res.stderr[-800:]

    wb = load_workbook(out)
    assert wb.sheetnames[:1] == ['About']
    for name in ('Values', 'Heads', 'Variants'):
        assert name in wb.sheetnames

    ws = wb['Values']
    hdr = [c.value for c in ws[2]]
    rows = {r[hdr.index('value')].value:
            dict(zip(hdr, [c.value for c in r]))
            for r in ws.iter_rows(min_row=3) if r[0].value}

    assert rows['Heavy Transport']['certainty'] == 'certain'
    # an alias is certain, and says where it goes
    assert rows['Transportation']['certainty'] == 'certain'
    assert rows['Transportation']['proposed official vertical'] \
        == 'Heavy Transport'
    # in use, in neither table: a decision
    assert rows['Heavy Cargo']['certainty'] == 'needs Nilesh'
    # An empty cell reads back as None, not '' — Excel has no concept
    # of an empty string distinct from an empty cell.
    assert not (rows['Heavy Cargo']['proposed official vertical'] or '')
    assert '(blank)' in rows
    assert rows['(blank)']['certainty'] == 'needs Nilesh'


def test_open_and_closed_are_counted_apart(master_world, tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / 'master.xlsx'
    _run('vertical_master_survey.py', '--out', str(out))
    ws = load_workbook(out)['Values']
    hdr = [c.value for c in ws[2]]
    for r in ws.iter_rows(min_row=3):
        row = dict(zip(hdr, [c.value for c in r]))
        if not row.get('value'):
            continue
        assert row['lead_total'] == row['lead_open'] + row['lead_closed']


def test_a_head_who_has_left_is_flagged(master_world, tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / 'master.xlsx'
    _run('vertical_master_survey.py', '--out', str(out))
    ws = load_workbook(out)['Heads']
    hdr = [c.value for c in ws[2]]
    rows = {r[hdr.index('vertical')].value:
            dict(zip(hdr, [c.value for c in r]))
            for r in ws.iter_rows(min_row=3) if r[0].value}
    assert rows['Heavy Transport']['status'] == 'HEAD HAS LEFT'
    assert rows['Management']['status'] == 'NO HEAD'


def test_variants_that_differ_only_by_spacing_are_grouped(master_world,
                                                          tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / 'master.xlsx'
    _run('vertical_master_survey.py', '--out', str(out))
    ws = load_workbook(out)['Variants']
    hdr = [c.value for c in ws[2]]
    found = [dict(zip(hdr, [c.value for c in r]))
             for r in ws.iter_rows(min_row=3) if r[0].value]
    collapsed = ' | '.join(f['values that collapse to it'] for f in found)
    assert 'Heavy Cargo' in collapsed and 'heavy  cargo' in collapsed


def test_the_decision_column_has_a_dropdown_allowing_blank(master_world,
                                                           tmp_path):
    import zipfile
    from openpyxl import load_workbook

    out = tmp_path / 'master.xlsx'
    _run('vertical_master_survey.py', '--out', str(out))
    wb = load_workbook(out)
    assert [c.value for c in wb['Values'][2]][-1] == 'Nilesh decision'
    with zipfile.ZipFile(out) as z:
        xml = ''.join(z.read(n).decode() for n in z.namelist()
                      if n.startswith('xl/worksheets/sheet'))
    i = xml.find('<dataValidation ')
    assert i != -1
    assert 'allowBlank="1"' in xml[i:i + 400]


def test_the_survey_changes_nothing(master_world, tmp_path):
    with flask_app.app_context():
        before = [(l.id, l.procam_vertical) for l in Lead.query.all()]
    _run('vertical_master_survey.py', '--out', str(tmp_path / 'm.xlsx'))
    with flask_app.app_context():
        after = [(l.id, l.procam_vertical) for l in Lead.query.all()]
    assert before == after


# ── CRM-VERTICAL-MASTER-01 Stage B ───────────────────────────────────
@pytest.fixture()
def master_b(world):
    """A lead on each mapped value, plus one awaiting review."""
    from app.models.email_admin import VerticalHead
    import scripts.vertical_master_apply as vm

    with flask_app.app_context():
        for code, name in (('DIR22010', 'James Francis Xavier'),
                           ('EMP372011', 'Sanjna Vardhan'),
                           ('DIR52011', 'Sethupathy Sundaram'),
                           ('DIR72012', 'Srinivas Marella'),
                           ('DIR12010', 'Nilesh Kumar Sinha')):
            _emp(code, name, role='user')
        if not VerticalHead.query.filter_by(vertical='All').first():
            db.session.add(VerticalHead(vertical='All', active=True))
        db.session.flush()

        made = {}
        for value in vm.MAPPING:
            lead = Lead(company=TAG + 'MB ' + value, assigned_to=ADMIN,
                        stage='Quoted', procam_vertical=value,
                        created_at=datetime.utcnow() - timedelta(days=3),
                        updated_at=datetime.utcnow())
            db.session.add(lead)
            db.session.flush()
            made[value] = lead.id
        # one sitting on Heavy Transport with keyword evidence: the
        # review population, which must not be touched
        waiting = Lead(company=TAG + 'MB waiting', assigned_to=ADMIN,
                       stage='Quoted', procam_vertical='Heavy Transport',
                       project='customs clearance, bill of entry',
                       created_at=datetime.utcnow() - timedelta(days=3),
                       updated_at=datetime.utcnow())
        db.session.add(waiting)
        db.session.flush()
        made['_waiting'] = waiting.id
        db.session.commit()
        yield made


def test_stage_b_dry_run_writes_nothing(master_b):
    import scripts.vertical_master_apply as vm

    res = _run('vertical_master_apply.py', '--check')
    assert res.returncode == 0, res.stderr[-900:]
    assert 'DRY RUN' in res.stdout
    with flask_app.app_context():
        for value in vm.MAPPING:
            assert db.session.get(
                Lead, master_b[value]).procam_vertical == value


def test_stage_b_moves_each_value_to_its_official_vertical(master_b):
    import scripts.vertical_master_apply as vm

    res = _run('vertical_master_apply.py', '--apply')
    assert res.returncode == 0, res.stderr[-900:]
    with flask_app.app_context():
        for value, target in vm.MAPPING.items():
            got = db.session.get(Lead, master_b[value]).procam_vertical
            assert got == target, f'{value} -> {got}, expected {target}'


def test_a_lead_awaiting_review_is_never_touched(master_b):
    _run('vertical_master_apply.py', '--apply')
    with flask_app.app_context():
        assert db.session.get(
            Lead, master_b['_waiting']).procam_vertical == 'Heavy Transport'


def test_every_moved_lead_carries_the_master_reason(master_b):
    from app.models.audit import AuditEvent
    import scripts.vertical_master_apply as vm

    _run('vertical_master_apply.py', '--apply')
    with flask_app.app_context():
        rows = AuditEvent.query.filter_by(
            action='lead.vertical_mastered').all()
        touched = {int(r.entity_id) for r in rows}
        for value in vm.MAPPING:
            assert master_b[value] in touched
        assert master_b['_waiting'] not in touched
        for r in rows:
            assert r.reason == 'vertical master, CRM-VERTICAL-MASTER-01'


def test_the_heads_are_set_and_all_is_removed(master_b):
    from app.models.email_admin import VerticalHead
    import scripts.vertical_master_apply as vm

    _run('vertical_master_apply.py', '--apply')
    with flask_app.app_context():
        rows = {v.vertical: v.head_emp_code
                for v in VerticalHead.query.all()}
        for vertical, code in vm.HEADS.items():
            assert rows.get(vertical) == code, vertical
        assert 'All' not in rows, 'an employee grouping, not a service'
        # Management is the fallback for anything unmapped and had
        # nobody; without it those notifications reach no one.
        assert rows['Management'] == 'DIR12010'


def test_the_dropdown_becomes_the_official_seven(master_b):
    import scripts.vertical_master_apply as vm

    _run('vertical_master_apply.py', '--apply')
    with flask_app.app_context():
        from app.services import vertical_audit as va
        va_list = va.acceptable_verticals(for_choice=True)
    for value in vm.MAPPING:
        assert value not in va_list, f'{value} is an alias now'
    for value in ('Heavy Transport', 'Freight Forwarding',
                  'General Transport', 'Project Freight'):
        assert value in va_list


def test_stage_b_is_idempotent(master_b):
    _run('vertical_master_apply.py', '--apply')
    with flask_app.app_context():
        first = {k: db.session.get(Lead, v).procam_vertical
                 for k, v in master_b.items()}
    res = _run('vertical_master_apply.py', '--apply')
    assert res.returncode == 0
    with flask_app.app_context():
        second = {k: db.session.get(Lead, v).procam_vertical
                  for k, v in master_b.items()}
    assert first == second


def test_blank_verticals_and_accounts_are_left_alone(master_b):
    with flask_app.app_context():
        blank = Lead(company=TAG + 'MB blank', assigned_to=ADMIN,
                     stage='Lost', procam_vertical=None,
                     created_at=datetime.utcnow() - timedelta(days=3),
                     updated_at=datetime.utcnow())
        db.session.add(blank)
        db.session.commit()
        blank_id = blank.id
        accounts_before = [(c.id, c.vertical) for c in Company.query.all()]

    _run('vertical_master_apply.py', '--apply')
    with flask_app.app_context():
        assert db.session.get(Lead, blank_id).procam_vertical is None
        assert [(c.id, c.vertical)
                for c in Company.query.all()] == accounts_before
