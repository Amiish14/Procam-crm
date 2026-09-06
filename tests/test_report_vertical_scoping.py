"""Reports must show a vertical head only their own vertical's data.

Builds two verticals with disjoint data, then asserts that each head sees
their own rows, never the other vertical's, and that an admin sees both.
"""
import os
import sys
import tempfile

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

os.environ['URL_PREFIX'] = ''
os.environ['SESSION_COOKIE_SECURE'] = 'false'
os.environ.setdefault('SECRET_KEY', 'test')
os.environ.setdefault('ADMIN_INITIAL_PASSWORD', 'ScopingTestOnly12345')
_DB = os.path.join(tempfile.mkdtemp(), 'scope.db')
os.environ['DATABASE_URL'] = 'sqlite:///' + _DB

import importlib                                          # noqa: E402
_main = importlib.import_module('app')
importlib.import_module('app.models')
flask_app, db = _main.app, _main.db

from app.models.task_engine import TaskInstance           # noqa: E402


HEAVY, PFM = 'Heavy Transport', 'PFM'


@pytest.fixture(scope='module')
def seeded():
    with flask_app.app_context():
        db.create_all()

        def emp(code, name, vertical, head=False, role='user'):
            e = _main.Employee.query.filter_by(emp_code=code).first()
            if not e:
                e = _main.Employee(emp_code=code, name=name)
                db.session.add(e)
            e.name, e.vertical, e.role = name, vertical, role
            e.is_vertical_head, e.is_active = head, True
            return e

        emp('ZADMIN', 'Scope Admin', 'Admin', role='admin')
        emp('ZHEAD_H', 'Head Heavy', HEAVY, head=True)
        emp('ZHEAD_P', 'Head PFM', PFM, head=True)
        emp('ZREP_H', 'Rep Heavy', HEAVY)
        emp('ZREP_P', 'Rep PFM', PFM)
        db.session.commit()

        for code, vert, tag in (('ZREP_H', HEAVY, 'H'), ('ZREP_P', PFM, 'P')):
            c = _main.Company(name=f'Acct-{tag}', pic_emp_code=code,
                              is_active=True)
            db.session.add(c)
            db.session.flush()
            lead = _main.Lead(company=f'Acct-{tag}', assigned_to=code,
                              procam_vertical=vert, source=f'src-{tag}')
            db.session.add(lead)
            db.session.flush()
            db.session.add(_main.LeadActivity(lead_id=lead.id, kind='call',
                                              subject=f'call-{tag}',
                                              performed_by=code))
            db.session.add(_main.Contact(name=f'Contact-{tag}',
                                         company=f'Acct-{tag}',
                                         assigned_to=code))
            db.session.add(_main.Opportunity(opp_number=f'OPP-{tag}-1',
                                             company_id=c.id,
                                             owner_emp_code=code,
                                             stage='Won', value_inr=1000))
            db.session.add(TaskInstance(task_key=f'lead.qualify.{tag}',
                                        entity_type='lead', entity_id=1,
                                        owner_user_id=code, status='Pending',
                                        priority=3))
        db.session.commit()
    return True


def _client_as(emp_code):
    c = flask_app.test_client()
    with flask_app.app_context():
        e = _main.Employee.query.filter_by(emp_code=emp_code).first()
        role, vertical = e.role, e.vertical or ''
    with c.session_transaction() as sess:
        sess['emp_code'] = emp_code
        sess['name'] = emp_code
        sess['role'] = role
        sess['vertical'] = vertical
    return c


def _body(emp_code, url):
    r = _client_as(emp_code).get(url)
    assert r.status_code == 200, f'{url} as {emp_code} → {r.status_code}'
    return r.get_data(as_text=True)


# Each report, and the marker proving whose data reached the page.
CASES = [
    ('/reports/tasks-by-user',          'ZREP_H',  'ZREP_P'),
    ('/reports/open-tasks',             'ZREP_H',  'ZREP_P'),
    ('/reports/accounts-by-pic',        'ZREP_H',  'ZREP_P'),
    ('/reports/contacts-developed',     'ZREP_H',  'ZREP_P'),
    ('/reports/won-value-by-account',   'Acct-H',  'Acct-P'),
    ('/reports/activities-by-account',  'Acct-H',  'Acct-P'),
    ('/reports/network-contribution',   'src-H',   'src-P'),
    ('/reports/dormant-accounts',       'Acct-H',  'Acct-P'),
    ('/reports/accounts-by-stage',      None,      None),
]


@pytest.mark.parametrize('url,mine,theirs', CASES)
def test_head_sees_only_own_vertical(seeded, url, mine, theirs):
    if mine is None:
        pytest.skip('aggregate has no per-vertical marker')
    body = _body('ZHEAD_H', url)
    assert mine in body, f'{url}: head lost their own data ({mine})'
    assert theirs not in body, f'{url}: LEAK — head saw {theirs}'


@pytest.mark.parametrize('url,mine,theirs', CASES)
def test_admin_sees_everything(seeded, url, mine, theirs):
    if mine is None:
        pytest.skip('aggregate has no per-vertical marker')
    body = _body('ZADMIN', url)
    assert mine in body and theirs in body, \
        f'{url}: admin should see both verticals'


def test_other_head_sees_the_mirror_image(seeded):
    body = _body('ZHEAD_P', '/reports/tasks-by-user')
    assert 'ZREP_P' in body
    assert 'ZREP_H' not in body


def test_plain_user_is_refused(seeded):
    r = _client_as('ZREP_H').get('/reports/tasks-by-user')
    assert r.status_code == 403


def test_hub_is_refused_to_plain_user(seeded):
    assert _client_as('ZREP_H').get('/reports').status_code == 403
    assert _client_as('ZHEAD_H').get('/reports').status_code == 200


def _xlsx_cells(resp):
    """Every cell value in the workbook, as strings.

    An .xlsx is a zip, so a substring check against the raw bytes can
    never fail — the workbook has to be opened for the check to mean
    anything.  (The first version of this test passed even with scoping
    disabled, for exactly that reason.)
    """
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(resp.get_data()), read_only=True)
    return [str(c.value) for ws in wb.worksheets
            for row in ws.iter_rows() for c in row if c.value is not None]


def test_xlsx_export_is_scoped_too(seeded):
    """A leak through ?format=xlsx would bypass every HTML check."""
    r = _client_as('ZHEAD_H').get('/reports/tasks-by-user?format=xlsx')
    assert r.status_code == 200
    cells = _xlsx_cells(r)
    assert 'ZREP_H' in cells, 'head lost their own data in the export'
    assert 'ZREP_P' not in cells, 'LEAK via xlsx export'


def test_xlsx_export_for_admin_has_both(seeded):
    cells = _xlsx_cells(_client_as('ZADMIN').get(
        '/reports/tasks-by-user?format=xlsx'))
    assert 'ZREP_H' in cells and 'ZREP_P' in cells
