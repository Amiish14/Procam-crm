"""Excel Master Data Import — §43-51.

The rules worth pinning: nothing is written without a confirmed preview
(§47), the mode is honoured (§48), an existing company is updated rather
than duplicated (§51), and an empty cell never erases a value someone
typed.
"""
import io
import os
import sys
import json
import tempfile

import pytest
from openpyxl import Workbook, load_workbook

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

os.environ['URL_PREFIX'] = ''
os.environ['SESSION_COOKIE_SECURE'] = 'false'
os.environ.setdefault('SECRET_KEY', 'test')
os.environ.setdefault('ADMIN_INITIAL_PASSWORD', 'XlTestOnly12345')
os.environ['DATABASE_URL'] = 'sqlite:///' + os.path.join(
    tempfile.mkdtemp(), 'xl.db')

import importlib                                          # noqa: E402
_main = importlib.import_module('app')
importlib.import_module('app.models')
flask_app, db = _main.app, _main.db
Company, Contact, Employee = _main.Company, _main.Contact, _main.Employee
flask_app.config['WTF_CSRF_ENABLED'] = False

from app.excel_io import service as xl                     # noqa: E402
from app.excel_io.templates import TEMPLATES, all_headers  # noqa: E402
from app.master_data import service as md                  # noqa: E402
from presales.models import AccountRelationshipTag         # noqa: E402


@pytest.fixture()
def clean():
    with flask_app.app_context():
        db.create_all()
        md.ensure_lists()
        for label in ('Customer', 'Vendor', 'Competitor'):
            md.add_item('relationship', label, label)
        md.add_item('industry', 'Power', 'Power')
        AccountRelationshipTag.query.delete()
        Contact.query.delete()
        Company.query.delete()
        e = Employee.query.filter_by(emp_code='XLADM').first()
        if not e:
            e = Employee(emp_code='XLADM', name='Xl Admin')
            db.session.add(e)
        e.role, e.is_active, e.must_change_pw = 'admin', True, False
        e.is_super_admin, e.vertical = True, 'All'
        db.session.commit()
    return True


def _sheet(kind, rows):
    """Build an upload the way a person would, from the real headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'DATA'
    ws.append(all_headers(kind))
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _records(kind, rows):
    return xl.parse(_sheet(kind, rows), kind)


# ── §43/§44: the template itself ─────────────────────────────────────
def test_template_has_data_instructions_and_lookups(clean):
    with flask_app.app_context():
        wb = load_workbook(xl.build_template('company'))
    assert 'DATA' in wb.sheetnames
    assert 'INSTRUCTIONS' in wb.sheetnames
    assert 'LOOKUPS' in wb.sheetnames


def test_lookups_come_from_master_data(clean):
    """§45 — a value added this morning is in the template this afternoon."""
    with flask_app.app_context():
        md.add_item('relationship', 'Freight Partner', 'Freight Partner')
        wb = load_workbook(xl.build_template('company'))
    values = {c.value for row in wb['LOOKUPS'].iter_rows() for c in row}
    assert 'Freight Partner' in values


# ── §46: validation ──────────────────────────────────────────────────
def test_missing_required_column_is_refused(clean):
    wb = Workbook()
    wb.active.title = 'DATA'
    wb.active.append(['Industry', 'City'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    records, problems = xl.parse(buf, 'company')
    assert problems and 'Company Name' in problems[0]


def test_unknown_lookup_value_is_an_error(clean):
    with flask_app.app_context():
        rows = [['Acme Ltd', 'Custmer', 'Power'] + [''] * 12]
        records, _p = _records('company', rows)
        results = xl.validate(records, 'company', xl.MODE_UPSERT)
    assert results[0]['errors']
    assert 'not a known relationship' in results[0]['errors'][0]


def test_unknown_employee_code_is_an_error(clean):
    with flask_app.app_context():
        rows = [['Acme Ltd', 'Customer', 'Power'] + [''] * 8
                + ['NOSUCHEMP', '', '', '']]
        records, _p = _records('company', rows)
        results = xl.validate(records, 'company', xl.MODE_UPSERT)
    assert any('not an employee code' in e
               for e in results[0]['errors'])


# ── §47: preview writes nothing ──────────────────────────────────────
def test_validation_writes_nothing(clean):
    with flask_app.app_context():
        before = Company.query.count()
        records, _p = _records('company', [['Brand New Co', 'Customer']
                                           + [''] * 13])
        xl.validate(records, 'company', xl.MODE_UPSERT)
        assert Company.query.count() == before


# ── §48: modes ───────────────────────────────────────────────────────
def test_create_only_skips_an_existing_company(clean):
    with flask_app.app_context():
        db.session.add(Company(name='Acme Limited', is_active=True))
        db.session.commit()
        records, _p = _records('company', [['Acme Ltd', 'Customer']
                                           + [''] * 13])
        results = xl.validate(records, 'company', xl.MODE_CREATE)
    assert results[0]['action'] == 'skip'


def test_update_only_skips_a_new_company(clean):
    with flask_app.app_context():
        records, _p = _records('company', [['Totally New Ltd', 'Customer']
                                           + [''] * 13])
        results = xl.validate(records, 'company', xl.MODE_UPDATE)
    assert results[0]['action'] == 'skip'


# ── §51: one company, however many files ─────────────────────────────
def test_the_same_company_is_updated_not_duplicated(clean):
    """Arriving as Customer in one file and Competitor in another must
    leave ONE record holding both."""
    with flask_app.app_context():
        from app import ImportBatch

        for relationship in ('Customer', 'Competitor'):
            records, _p = _records(
                'company', [['Vedanta Limited', relationship, 'Power']
                            + [''] * 12])
            results = xl.validate(records, 'company', xl.MODE_UPSERT)
            batch = ImportBatch(kind='company', filename='t.xlsx')
            db.session.add(batch)
            db.session.commit()
            xl.commit(results, 'company', xl.MODE_UPSERT, batch, 'XLADM')

        matches = Company.query.filter(
            Company.name.ilike('%vedanta%')).all()
        assert len(matches) == 1, \
            f'{len(matches)} Vedanta records — the import duplicated it'

        tags = {t.tag for t in AccountRelationshipTag.query.filter_by(
            account_id=matches[0].id).all()}
        assert tags == {'Customer', 'Competitor'}, \
            'classifications should accumulate on the one record'


def test_an_empty_cell_never_erases_existing_data(clean):
    """An import that blanks values people typed is worse than one that
    skips them."""
    with flask_app.app_context():
        from app import ImportBatch
        db.session.add(Company(name='Keepsake Ltd', is_active=True,
                               city='Mumbai', industry='Power'))
        db.session.commit()

        records, _p = _records('company', [['Keepsake Ltd', 'Customer']
                                           + [''] * 13])
        results = xl.validate(records, 'company', xl.MODE_UPSERT)
        batch = ImportBatch(kind='company', filename='t.xlsx')
        db.session.add(batch)
        db.session.commit()
        xl.commit(results, 'company', xl.MODE_UPSERT, batch, 'XLADM')

        kept = Company.query.filter_by(name='Keepsake Ltd').first()
        assert kept.city == 'Mumbai', 'an empty cell erased a real value'
        assert kept.industry == 'Power'


# ── §49: the error report ────────────────────────────────────────────
def test_error_report_lists_row_problem_and_fix(clean):
    with flask_app.app_context():
        records, _p = _records('company', [
            ['Good Co', 'Customer'] + [''] * 13,
            ['', 'Customer'] + [''] * 13,
        ])
        results = xl.validate(records, 'company', xl.MODE_UPSERT)
        wb = load_workbook(xl.build_error_report(results, 'company'))
    ws = wb['ERRORS']
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert body, 'the failing row is missing from the report'
    assert body[0][0] == 3, 'the row number must match the source file'
    assert 'required' in str(body[0][2]).lower()
    assert body[0][3], 'every error needs a suggested correction'


def test_summary_counts_reconcile(clean):
    with flask_app.app_context():
        records, _p = _records('company', [
            ['Alpha Co', 'Customer'] + [''] * 13,
            ['', 'Customer'] + [''] * 13,
            ['Beta Co', 'Nonsense'] + [''] * 13,
        ])
        results = xl.validate(records, 'company', xl.MODE_UPSERT)
        s = xl.summarise(results)
    assert s['total'] == 3
    assert s['errors'] == 2
    assert s['valid'] == 1


def test_page_and_template_need_admin(clean):
    anon = flask_app.test_client()
    assert anon.get('/admin/import').status_code in (302, 401, 403)
    assert anon.get('/admin/import/template/company'
                    ).status_code in (302, 401, 403)


# ── Overseas agents (§43) ────────────────────────────────────────────
@pytest.fixture()
def agents_ready(clean):
    with flask_app.app_context():
        for label in ('Overseas Agent', 'Overseas Partner'):
            md.add_item('relationship', label, label)
        for net in ('PCN', 'WCA', 'THLG'):
            md.add_item('network', net, net)
    return True


def test_overseas_agent_template_exists_and_is_complete(agents_ready):
    with flask_app.app_context():
        wb = load_workbook(xl.build_template('overseas_agent'))
    assert set(wb.sheetnames) == {'DATA', 'INSTRUCTIONS', 'LOOKUPS'}
    headers = [c.value for c in wb['DATA'][1]]
    for expected in ('Company Name', 'Relationship', 'Country', 'Networks',
                     'Contact Name', 'Contact Email'):
        assert expected in headers, f'{expected} missing from the template'


def test_agent_template_lookups_carry_networks(agents_ready):
    with flask_app.app_context():
        wb = load_workbook(xl.build_template('overseas_agent'))
    values = {c.value for row in wb['LOOKUPS'].iter_rows() for c in row}
    assert 'PCN' in values
    assert 'Overseas Agent' in values


def _agent_row(name, relationship='Overseas Agent', country='United States',
               networks='', person='', email=''):
    """A row in the template's own column order."""
    return [name, relationship, country, 'New York', '', '', '', '',
            networks, person, '', email, '', '', '']


def test_uploading_agents_maps_to_company_master(agents_ready):
    """The whole point: fill it in, upload it, and it lands correctly."""
    with flask_app.app_context():
        from app import ImportBatch
        records, problems = _records('overseas_agent', [
            _agent_row('Logfret Inc.', 'Overseas Agent', 'United States',
                       'PCN, WCA', 'Maria Silva', 'maria@logfret.com'),
        ])
        assert not problems, problems
        results = xl.validate(records, 'overseas_agent', xl.MODE_UPSERT)
        assert not results[0]['errors'], results[0]['errors']

        batch = ImportBatch(kind='overseas_agent', filename='a.xlsx')
        db.session.add(batch)
        db.session.commit()
        counts = xl.commit(results, 'overseas_agent', xl.MODE_UPSERT,
                           batch, 'XLADM')
        assert counts['created'] == 1

        company = Company.query.filter(
            Company.name.ilike('%logfret%')).first()
        assert company is not None
        assert company.country == 'United States'

        tags = {t.tag for t in AccountRelationshipTag.query.filter_by(
            account_id=company.id).all()}
        assert 'Overseas Agent' in tags, 'the classification was not applied'
        assert {'PCN', 'WCA'} <= tags, \
            'network memberships were dropped — relationship and networks '\
            'must BOTH be applied, not either'

        contact = Contact.query.filter_by(name='Maria Silva').first()
        assert contact is not None
        assert contact.company_id == company.id


def test_country_is_required_for_an_agent(agents_ready):
    with flask_app.app_context():
        records, _p = _records('overseas_agent',
                               [_agent_row('No Country Ltd', country='')])
        results = xl.validate(records, 'overseas_agent', xl.MODE_UPSERT)
    assert any('Country is required' in e for e in results[0]['errors'])


def test_an_agent_that_already_exists_is_updated_not_duplicated(agents_ready):
    """An agent already on file as a customer gains the classification."""
    with flask_app.app_context():
        from app import ImportBatch
        db.session.add(Company(name='Dual Role Logistics Ltd',
                               is_active=True))
        db.session.commit()

        records, _p = _records('overseas_agent', [
            _agent_row('Dual Role Logistics', 'Overseas Partner', 'Germany',
                       'THLG')])
        results = xl.validate(records, 'overseas_agent', xl.MODE_UPSERT)
        batch = ImportBatch(kind='overseas_agent', filename='a.xlsx')
        db.session.add(batch)
        db.session.commit()
        xl.commit(results, 'overseas_agent', xl.MODE_UPSERT, batch, 'XLADM')

        matches = Company.query.filter(
            Company.name.ilike('%dual role%')).all()
        assert len(matches) == 1, 'the agent was duplicated'
        tags = {t.tag for t in AccountRelationshipTag.query.filter_by(
            account_id=matches[0].id).all()}
        assert 'Overseas Partner' in tags and 'THLG' in tags


def test_several_rows_for_one_agent_make_one_company(agents_ready):
    with flask_app.app_context():
        from app import ImportBatch
        records, _p = _records('overseas_agent', [
            _agent_row('Multi Contact Freight', person='Ann One',
                       email='ann@mcf.com'),
            _agent_row('Multi Contact Freight', person='Bob Two',
                       email='bob@mcf.com'),
        ])
        results = xl.validate(records, 'overseas_agent', xl.MODE_UPSERT)
        batch = ImportBatch(kind='overseas_agent', filename='a.xlsx')
        db.session.add(batch)
        db.session.commit()
        xl.commit(results, 'overseas_agent', xl.MODE_UPSERT, batch, 'XLADM')

        assert Company.query.filter(
            Company.name.ilike('%multi contact%')).count() == 1
        assert Contact.query.filter(
            Contact.email.in_(['ann@mcf.com', 'bob@mcf.com'])).count() == 2
