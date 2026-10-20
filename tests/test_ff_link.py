"""
FF-LINK-01 — a shared email domain is not a shared company.

The vendor masking finds a client's contacts through the account
link, so an unlinked lead can carry a buyer's name and mobile to a
carrier. Linking them is therefore a prerequisite for turning FF on,
and the temptation is to take every domain match.

Production says otherwise. `renew.com` carries both RENEW POWER and
RENEW PHOTOVOLTAICS; accounts exist called "Mail" and "Asia",
auto-created from `mail.oyaso.com` and `asia.meap.com`, which match
by domain perfectly while naming nobody. 71 matches became 51 to
write and 20 for a person once the names had to agree as well.
"""
import os
import sys
import tempfile

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

os.environ['URL_PREFIX'] = ''
os.environ.setdefault('SECRET_KEY', 'fflink-secret')
os.environ.setdefault('ADMIN_INITIAL_PASSWORD', 'FfLinkTest12345')
os.environ.setdefault(
    'DATABASE_URL',
    'sqlite:///' + os.path.join(tempfile.mkdtemp(), 'fflink.db'))

import importlib                                          # noqa: E402
_main = importlib.import_module('app')
importlib.import_module('app.models')
flask_app = _main.app

import scripts.ff_link_accounts as link                   # noqa: E402


@pytest.mark.parametrize('lead_name,account_name', [
    ('BEML LIMITED', 'BEML LIMITED'),
    ('AGUNSA', 'AGUNSA'),
    ('SKIPPER SEIL LTD', 'Skipper Seil Ltd'),
    ('TEMA INDIA LIMITED', 'TEMA INDIA LIMITED'),
])
def test_the_same_company_agrees(lead_name, account_name):
    assert link._names_agree(lead_name, account_name)


@pytest.mark.parametrize('lead_name,account_name,why', [
    ('Oyaso', 'Mail', 'an account auto-named from a subdomain'),
    ('Meap', 'Asia', 'an account auto-named from a subdomain'),
    ('RENEW PHOTOVOLTAICS PRIVATE LIMITED', 'RENEW POWER PRIVATE LIMITED',
     'two legal entities on one corporate domain'),
    ('Vikram Solar Limited', 'VSL Green Power Pvt Ltd',
     'a different name entirely'),
])
def test_a_shared_domain_is_not_a_shared_company(lead_name,
                                                 account_name, why):
    assert not link._names_agree(lead_name, account_name), why


def test_an_empty_name_never_agrees():
    assert not link._names_agree('', 'Anything')
    assert not link._names_agree('Something', '')
    assert not link._names_agree(None, None)


def test_a_very_short_overlap_does_not_count():
    """Two letters inside a longer name is a coincidence."""
    assert not link._names_agree('AB', 'ABC Logistics Limited')


def test_only_trusted_rules_and_agreeing_names_are_written():
    """The rules a person set, or the exact address on a contact."""
    assert 'contact email' in link.TRUSTED_RULES
    assert 'learned email' in link.TRUSTED_RULES
    assert 'company name' in link.TRUSTED_RULES
    assert 'account email domain' not in link.TRUSTED_RULES, (
        'a domain match must earn its place by the names agreeing')


def test_the_buckets_are_four_not_two():
    """`resolve` returns nothing both for "no candidate" and "several",
    which is right for it and useless for a report."""
    assert {link.MATCHED, link.REVIEW, link.AMBIGUOUS, link.NONE} == {
        'matched', 'needs a look', 'ambiguous', 'no match'}


def test_unprotectable_names_are_asked_of_the_masking_rule():
    """Not a word list kept beside it: `variants()` is what decides
    whether a name can be redacted, so it is what decides whether an
    alias is needed."""
    from app.services import ff_privacy
    # Too short to redact safely — these are the ones needing an
    # alias, and production has exactly four: Acs, Efl, Qd, Uht.
    assert ff_privacy.variants('Acs') == set()
    assert ff_privacy.variants('Qd') == set()
    # A whole name that happens to be a generic word IS protected:
    # the full name is always a variant, and the generic list only
    # governs single tokens drawn out of a longer name. The cost is
    # that "steel coils" gets masked too, which is the right way
    # round — leaking the client is worse than reading badly.
    assert ff_privacy.variants('Steel') == {'Steel'}
    # And in a longer name, the generic token alone is not a secret
    # while the distinctive run is.
    v = ff_privacy.variants('Tata Steel Limited')
    assert 'Tata Steel' in v and 'Tata' in v
    assert 'Steel' not in v


def test_export_review_actually_writes_a_sheet():
    """`--export-review` raised NameError in production.

    `_write_review` was defined below the `if __name__` block, so the
    module called `main()` while the name was still unbound. Every
    test here imported the module and exercised `survey()`, which
    proved the matching and never once ran the script as a script.
    So this one runs it the way a person does, in a subprocess, and
    opens what comes out.
    """
    import subprocess
    out = subprocess.run(
        [sys.executable, os.path.join(_ROOT, 'scripts',
                                      'ff_link_accounts.py'),
         '--export-review'],
        capture_output=True, text=True, cwd=_ROOT,
        env={**os.environ, 'PYTHONPATH': _ROOT})
    assert out.returncode == 0, out.stderr[-2000:]
    path = [ln.strip() for ln in out.stdout.splitlines()
            if ln.strip().endswith('.xlsx')]
    assert path, f'no workbook named in the output:\n{out.stdout[-2000:]}'
    from openpyxl import load_workbook
    ws = load_workbook(path[-1]).active
    header = [c.value for c in ws[4]]
    assert 'correct account' in [str(h or '').lower() for h in header], \
        f'the column a person fills in is missing: {header}'
    os.remove(path[-1])
