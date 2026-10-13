"""
Whether an employee's address can actually receive CRM mail.

Three quarters of the CRM's undeliverable email came from the employee
master rather than from the mail server. Three kinds of bad address,
all of them put there by a person in a hurry:

    an address on a domain the CRM is not allowed to write to — the
    mail policy blocks it, correctly, and the notification silently
    reaches nobody;

    a placeholder minted from the employee code, `emp101@…`,
    `dir12010@…`, `con77@…` — created so a form would save, never a
    real mailbox, and the server answers 5.1.1;

    the same address on two employees, usually a shared departmental
    mailbox, which means a notification meant for one person is read
    by whoever gets there first and dismissed.

None of these are mail problems and none of them are visible from the
Email Health screen, because a message the CRM never sent does not
appear in the log. So they are caught here, where the address is
entered, and listed on the health card for the ones already on file.

    check(address, emp_code=..., ignore_id=...)
        → {'ok', 'errors': [...], 'warnings': [...], 'normalised'}

Errors refuse the save. Warnings do not — a shared mailbox is
sometimes exactly what a department wants, and the CRM is not in a
position to overrule that. It just says so.
"""
from __future__ import annotations

import logging
import re

log = logging.getLogger(__name__)

#: A placeholder minted from an employee code rather than a mailbox.
#: The three prefixes the CRM's own employee codes use, and nothing
#: else — `operations@` is a real mailbox and must not be refused.
PLACEHOLDER = re.compile(r'^(emp|dir|con)\d+@', re.I)

#: Deliberately permissive. The authority on whether an address exists
#: is the mail server, not a regular expression; this catches the
#: typos that stop it being an address at all.
SHAPE = re.compile(r'^[^@\s,;<>"]+@[^@\s,;<>"]+\.[A-Za-z]{2,}$')


def normalise(address):
    """Lower-cased, trimmed, display name stripped."""
    text = (address or '').strip()
    if '<' in text and '>' in text:
        text = text[text.rfind('<') + 1:text.rfind('>')].strip()
    return text.lower()


def allowed_domains():
    """Every domain an employee address may use.

    The mail policy's list — the domains the CRM is permitted to send
    to — plus anything an administrator has added in the settings
    screen. The two are the same question asked twice, so they share
    an answer: an address the CRM may not send to is not a usable
    employee address.
    """
    from app.services import mail_policy

    out = list(mail_policy.internal_domains())
    try:
        from app.services import admin_settings
        for extra in (admin_settings.get('allowed_email_domains') or []):
            dom = str(extra).strip().lower().lstrip('@')
            if dom and dom not in out:
                out.append(dom)
    except Exception:
        pass
    return out


def domain_ok(address):
    dom = (normalise(address).rsplit('@', 1) + [''])[1]
    if not dom:
        return False
    for allowed in allowed_domains():
        if dom == allowed or dom.endswith('.' + allowed):
            return True
    return False


def duplicate_holder(address, *, ignore_id=None):
    """The other employee already using this address, or None."""
    from app import Employee

    addr = normalise(address)
    if not addr:
        return None
    try:
        query = Employee.query.filter(
            db_lower(Employee.email) == addr)
        if ignore_id is not None:
            query = query.filter(Employee.id != ignore_id)
        return query.first()
    except Exception:
        return None


def db_lower(column):
    from sqlalchemy import func
    return func.lower(func.trim(column))


def check(address, *, emp_code=None, ignore_id=None, required=True):
    """Is this address usable? See the module docstring."""
    out = {'ok': True, 'errors': [], 'warnings': [],
           'normalised': normalise(address)}
    addr = out['normalised']

    if not addr:
        if required:
            out['errors'].append('An email address is required — without '
                                 'one this person receives no CRM email.')
        out['ok'] = not out['errors']
        return out

    if not SHAPE.match(addr):
        out['errors'].append(f'{addr} is not a valid email address.')
    elif PLACEHOLDER.match(addr):
        out['errors'].append(
            f'{addr} looks like a placeholder built from an employee code, '
            f'not a real mailbox. Mail to it is refused by the server and '
            f'the person is told nothing.')
    elif not domain_ok(addr):
        out['errors'].append(
            f'The CRM is not allowed to send to @{addr.rsplit("@", 1)[1]}. '
            f'Allowed: ' + ', '.join('@' + d for d in allowed_domains()) +
            '. An address outside these is blocked before it is sent.')

    other = duplicate_holder(addr, ignore_id=ignore_id)
    if other is not None and (emp_code or '').upper() != (
            other.emp_code or '').upper():
        out['warnings'].append(
            f'{addr} is already on {other.emp_code} '
            f'({(other.name or "").strip() or "no name on file"}). Both '
            f'people will receive each other\'s notifications, and whoever '
            f'reads one first dismisses it for the other.')

    out['ok'] = not out['errors']
    return out


# ── the health card ──────────────────────────────────────────────────
def unreachable(limit=200):
    """Active employees who cannot receive CRM mail, and why.

    [{'emp_code', 'name', 'email', 'vertical', 'reason'}]. The reason
    is the first thing wrong with it, which is the thing to fix.
    """
    from app import Employee

    out = []
    try:
        rows = (Employee.query
                .filter(Employee.is_active.is_(True))
                .order_by(Employee.emp_code).all())
    except Exception:
        return out

    seen = {}
    for emp in rows:
        addr = normalise(getattr(emp, 'email', ''))
        reason = ''
        if not addr:
            reason = 'no email address on file'
        elif not SHAPE.match(addr):
            reason = 'not a valid email address'
        elif PLACEHOLDER.match(addr):
            reason = 'a placeholder built from the employee code'
        elif not domain_ok(addr):
            reason = f'@{addr.rsplit("@", 1)[1]} is not an allowed domain'
        elif addr in seen:
            reason = f'the same address as {seen[addr]}'
        if addr and not reason:
            seen[addr] = emp.emp_code
        if reason:
            out.append({'emp_code': emp.emp_code,
                        'name': (emp.name or '').strip(),
                        'email': addr,
                        'vertical': (getattr(emp, 'vertical', '') or '').strip(),
                        'reason': reason})
        if len(out) >= limit:
            break
    return out


# ── the HRMS import ──────────────────────────────────────────────────
#: Column headings the HRMS export uses, and the ones people rename
#: them to. Matched case- and space-insensitively.
HEADINGS = {
    'emp_code': ('emp code', 'employee code', 'empcode', 'code', 'emp id',
                 'employee id'),
    'name': ('name', 'employee name', 'full name'),
    'email': ('email', 'email id', 'e-mail', 'official email',
              'email address', 'company email'),
    'vertical': ('vertical', 'department', 'business vertical', 'division'),
    'designation': ('designation', 'title', 'role'),
}


def _heading_map(header):
    """{column index: field} for a header row."""
    out = {}
    for idx, cell in enumerate(header or []):
        key = ' '.join(str(cell or '').strip().lower().replace('_', ' ')
                       .split())
        for field, names in HEADINGS.items():
            if key in names and field not in out.values():
                out[idx] = field
                break
    return out


def read_rows(path):
    """(rows, problem). Rows are dicts keyed by the fields above.

    Reads CSV and XLSX, because the HRMS exports one and people send
    the other. Nothing is written: this is the preview step, and an
    import nobody has looked at is how the employee master got the
    placeholder addresses in the first place.
    """
    name = (path or '').lower()
    try:
        if name.endswith(('.xlsx', '.xlsm')):
            table = _read_xlsx(path)
        else:
            table = _read_csv(path)
    except Exception as exc:                                  # noqa: BLE001
        log.exception('could not read %s', path)
        return [], f'Could not read the file: {exc}'

    table = [r for r in table if any(str(c or '').strip() for c in r)]
    if not table:
        return [], 'The file is empty.'

    headings = _heading_map(table[0])
    if 'email' not in headings.values():
        return [], ('No email column found. The file needs a heading of '
                    'Email, Email ID or Official Email.')
    if 'emp_code' not in headings.values():
        return [], ('No employee code column found. The file needs a '
                    'heading of Emp Code or Employee Code — the CRM '
                    'matches people on the code, never on the name.')

    rows = []
    for raw in table[1:]:
        row = {}
        for idx, field in headings.items():
            if idx < len(raw):
                row[field] = str(raw[idx] or '').strip()
        if not (row.get('emp_code') or '').strip():
            continue
        row['emp_code'] = row['emp_code'].strip().upper()
        row['email'] = normalise(row.get('email'))
        rows.append(row)
    return rows, ''


def _read_csv(path):
    import csv

    for encoding in ('utf-8-sig', 'utf-16', 'latin-1'):
        try:
            with open(path, newline='', encoding=encoding) as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
                except Exception:
                    dialect = csv.excel
                return [list(r) for r in csv.reader(handle, dialect)]
        except UnicodeDecodeError:
            continue
    raise ValueError('the file is not text the CRM can read')


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:                                       # pragma: no cover
        raise ValueError(
            'XLSX needs the openpyxl package, which is not installed. '
            'Save the sheet as CSV and import that instead.')
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def preview(path):
    """What the import would do, before it does any of it.

    Returns (plan, problem). `plan` is a list of
    {'row', 'emp_code', 'name', 'email', 'action', 'detail'} where
    action is one of: update, unchanged, not_found, invalid.

    No employee is created. An HRMS export covers the whole company
    and the CRM covers the people who use it; creating a record for
    every name in the file would add several hundred employees nobody
    asked for, so a code the CRM does not know is reported rather
    than added.
    """
    from app import Employee

    rows, problem = read_rows(path)
    if problem:
        return [], problem

    plan = []
    for number, row in enumerate(rows, start=2):
        code = row['emp_code']
        emp = Employee.query.filter_by(emp_code=code).first()
        item = {'row': number, 'emp_code': code,
                'name': row.get('name', ''), 'email': row.get('email', ''),
                'vertical': row.get('vertical', ''),
                'action': '', 'detail': ''}
        if emp is None:
            item['action'] = 'not_found'
            item['detail'] = ('no employee with this code — the CRM will '
                              'not create one from an HRMS export')
            plan.append(item)
            continue
        verdict = check(row.get('email'), emp_code=code, ignore_id=emp.id)
        if not verdict['ok']:
            item['action'] = 'invalid'
            item['detail'] = ' '.join(verdict['errors'])
            plan.append(item)
            continue
        current = normalise(getattr(emp, 'email', ''))
        if current == item['email']:
            item['action'] = 'unchanged'
            item['detail'] = 'already on file'
        else:
            item['action'] = 'update'
            item['detail'] = (f'{current or "(blank)"} → {item["email"]}'
                              + ('. ' + ' '.join(verdict['warnings'])
                                 if verdict['warnings'] else ''))
        plan.append(item)
    return plan, ''


def apply_plan(plan, *, actor=None):
    """Write the `update` rows from a preview. Returns a small report.

    Only `update` rows, and only the email address. An HRMS export is
    authoritative about addresses; it is not authoritative about who
    owns which lead, which vertical the CRM files them under, or
    whether they are still active here.
    """
    from app import Employee, db

    report = {'updated': 0, 'skipped': 0, 'errors': []}
    for item in (plan or []):
        if item.get('action') != 'update':
            report['skipped'] += 1
            continue
        emp = Employee.query.filter_by(emp_code=item['emp_code']).first()
        if emp is None:
            report['errors'].append(f'{item["emp_code"]}: disappeared '
                                    f'between the preview and the import')
            continue
        was = getattr(emp, 'email', '')
        emp.email = item['email']
        report['updated'] += 1
        try:
            from app.services import audit
            audit.record('employee.email_changed', 'Employee',
                         emp.emp_code, old={'email': was},
                         new={'email': item['email']}, actor=actor,
                         reason='HRMS import')
        except Exception:
            pass
    try:
        db.session.commit()
    except Exception as exc:                                  # noqa: BLE001
        db.session.rollback()
        report['errors'].append(f'nothing was saved: {exc}')
        report['updated'] = 0
    return report
