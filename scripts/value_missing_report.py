"""
Open leads past RFQ with no value, per owner and stage. Read-only.

    .venv/bin/python scripts/value_missing_report.py

Nothing is filled in. The whole point of CRM-VALUE-01 is that the
number comes from the person who knows it; a script that guesses one
produces a forecast built on arithmetic nobody stands behind, which
is worse than the gap it covers.

Writes docs/exports/value_missing_<date>.xlsx and prints the counts.
"""
import argparse
import os
import sys
from collections import Counter, defaultdict
from datetime import date, datetime

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _ROOT)

from app import Employee, Lead, app, db                   # noqa: E402
from app.services import lead_value, value_gate           # noqa: E402


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out')
    args = ap.parse_args()

    with app.app_context():
        today = date.today()
        print(f'database: {db.engine.url}')
        print(f'as of:    {datetime.utcnow()} UTC  (read-only)\n')

        names = {e.emp_code: e.name for e in Employee.query.all()}
        rows = (Lead.query
                .filter(Lead.is_archived.isnot(True),
                        Lead.stage.in_(value_gate.OPEN_REQUIRES_VALUE))
                .order_by(Lead.assigned_to.asc(), Lead.stage.asc()).all())
        missing = [l for l in rows if lead_value.value_inr(l) is None]

        grid = defaultdict(Counter)
        for l in missing:
            grid[l.assigned_to or '(unowned)'][l.stage or '(none)'] += 1

        stages = [s for s in value_gate.OPEN_REQUIRES_VALUE
                  if any(c.get(s) for c in grid.values())]
        head = f'{"owner":<14}{"name":<26}' + ''.join(
            f'{s[:17]:>19}' for s in stages) + f'{"total":>8}'
        print(head)
        print('-' * len(head))
        for owner in sorted(grid, key=lambda o: -sum(grid[o].values())):
            counts = grid[owner]
            print(f'{owner:<14}{(names.get(owner, "") or "")[:25]:<26}'
                  + ''.join(f'{counts.get(s, 0):>19}' for s in stages)
                  + f'{sum(counts.values()):>8}')
        print('-' * len(head))
        totals = Counter()
        for c in grid.values():
            totals.update(c)
        print(f'{"TOTAL":<14}{"":<26}'
              + ''.join(f'{totals.get(s, 0):>19}' for s in stages)
              + f'{len(missing):>8}')

        print(f'\n  {len(missing)} of {len(rows)} lead(s) at RFQ Generated '
              f'or later carry no value.')
        print(f'  {len(grid)} owner(s) affected.')
        print('\n  Nothing was filled in. These need the person who '
              'knows the number.')

        out = args.out or os.path.join(
            _ROOT, 'docs', 'exports',
            f'value_missing_{today.isoformat()}.xlsx')
        _write(missing, names, out, stages, grid)
        print(f'\n  {out}')


def _write(missing, names, path, stages, grid):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'By owner'
    ws.append(['Open leads past RFQ with no value'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f'{len(missing)} lead(s). Nothing is filled in by the '
               f'CRM — each needs the person who knows the number.'])
    ws.append([])
    ws.append(['Owner', 'Name'] + list(stages) + ['Total'])
    for cell in ws[4]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F2937')
    for owner in sorted(grid, key=lambda o: -sum(grid[o].values())):
        c = grid[owner]
        ws.append([owner, names.get(owner, '')]
                  + [c.get(s, 0) for s in stages] + [sum(c.values())])
    for col, w in [('A', 14), ('B', 26)]:
        ws.column_dimensions[col].width = w

    det = wb.create_sheet('Every lead')
    det.append(['Lead', 'Company', 'Project', 'Stage', 'Owner', 'Name',
                'Vertical', 'Created', 'Quote date', 'Decision (yours)'])
    for cell in det[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F2937')
    for l in missing:
        det.append([l.id, l.company or '', l.project or '', l.stage or '',
                    l.assigned_to or '', names.get(l.assigned_to, ''),
                    l.procam_vertical or '',
                    str(l.created_at)[:10] if l.created_at else '',
                    str(l.quote_date) if l.quote_date else '', ''])
    for i, w in enumerate((8, 34, 30, 20, 12, 24, 20, 12, 12, 22), start=1):
        det.column_dimensions[det.cell(1, i).column_letter].width = w
    det.freeze_panes = 'A2'
    det.auto_filter.ref = f'A1:J{det.max_row}'
    wb.save(path)


if __name__ == '__main__':
    main()
