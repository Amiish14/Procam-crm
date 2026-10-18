"""
CRM-VERTICAL-FIX-01 Stage A — the repair workbook.

READ-ONLY. Writes one .xlsx and changes no lead.

    .venv/bin/python scripts/vertical_repair_workbook.py
    .venv/bin/python scripts/vertical_repair_workbook.py --since 2026-09-01

Three sheets, because the suspected leads are not one population and
treating them as one is how a repair becomes a second incident. See
`app/services/vertical_audit.py` for what separates them.

  Recorded       the audit trail holds the previous value, so the
                 restore is known rather than inferred. Stage B
                 applies this sheet and only this sheet.
  Review         keyword evidence only, with a blank "correct
                 vertical" column for a person to fill in.
  Changed since  somebody set the vertical again after the rewrite.
                 Never touched.
"""
import argparse
import os
import sys
from datetime import date

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))

from app import app, db                                   # noqa: E402
from app.services import vertical_audit as va             # noqa: E402

SHEETS = (
    ('Recorded', 'recorded',
     ('lead_id', 'company', 'title', 'stage', 'pic', 'now',
      'restore_to', 'rewritten_by', 'rewritten_at', 'looks_like',
      'agrees', 'evidence', 'updated_at')),
    ('Review', 'review',
     ('lead_id', 'company', 'title', 'stage', 'pic', 'now',
      'looks_like', 'confidence', 'keyword_hits', 'heavy_hits',
      'evidence', 'correct vertical', 'updated_at')),
    ('Changed since', 'changed_since',
     ('lead_id', 'company', 'title', 'stage', 'pic', 'now',
      'was', 'rewritten_by', 'rewritten_at',
      'changed_to', 'changed_by', 'changed_at', 'evidence')),
)

NOTES = {
    'Recorded': (
        'The audit trail records this rewrite: who did it, when, and '
        'what the vertical was before. "restore_to" is that previous '
        'value — it is not a guess. Stage B applies this sheet in one '
        'transaction, with an audit row per lead.'),
    'Review': (
        'Keyword evidence only. There is no audit row, so the CRM does '
        'not know what the vertical was — only that the lead\'s own '
        'words disagree with Heavy Transport. Fill in "correct '
        'vertical" where you are sure and leave it blank where you are '
        'not; blanks are skipped.'),
    'Changed since': (
        'Somebody set the vertical again AFTER the rewrite. Restoring '
        'would undo a decision taken since. These are listed so they '
        'are accounted for, and are never written to.'),
}


def _write(path, data, since):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill('solid', fgColor='BC1D2F')
    head_font = Font(color='FFFFFF', bold=True, size=10)
    note_font = Font(italic=True, size=9, color='6B6762')

    for title, key, columns in SHEETS:
        rows = data[key]
        ws = wb.create_sheet(title)
        ws['A1'] = NOTES[title]
        ws['A1'].font = note_font
        ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=max(4, len(columns)))
        ws.row_dimensions[1].height = 42

        for col, name in enumerate(columns, start=1):
            cell = ws.cell(row=2, column=col, value=name)
            cell.fill, cell.font = head_fill, head_font
        for r, row in enumerate(rows, start=3):
            for col, name in enumerate(columns, start=1):
                ws.cell(row=r, column=col, value=row.get(name, ''))

        widths = {'company': 34, 'title': 40, 'evidence': 38,
                  'looks_like': 18, 'restore_to': 18, 'changed_to': 18,
                  'correct vertical': 18, 'was': 18, 'agrees': 12}
        for col, name in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col)].width = \
                widths.get(name, 14)
        ws.freeze_panes = 'A3'
        if rows:
            ws.auto_filter.ref = (f'A2:{get_column_letter(len(columns))}'
                                  f'{len(rows) + 2}')

    summary = wb.create_sheet('About', 0)
    summary['A1'] = 'CRM-VERTICAL-FIX-01 — vertical repair'
    summary['A1'].font = Font(bold=True, size=13)
    lines = [
        '',
        f'Generated {date.today()} from the live database. Read-only: '
        f'no lead was changed.',
        f'Window: leads on Heavy Transport updated since {since}.',
        f'Examined: {data["examined"]}',
        '',
        f'Recorded      {len(data["recorded"]):>5}   restore is known '
        f'from the audit trail — Stage B applies these',
        f'Review        {len(data["review"]):>5}   keyword evidence '
        f'only — needs a person',
        f'Changed since {len(data["changed_since"]):>5}   set again '
        f'after the rewrite — never touched',
        '',
        'The lead edit form offered four verticals against the eleven '
        'leads carry. A select with no option matching the record '
        'selects option zero, so saving anything at all rewrote the '
        'vertical to Heavy Transport. The form is fixed; this is the '
        'damage already done.',
    ]
    for i, line in enumerate(lines, start=2):
        summary[f'A{i}'] = line
    summary.column_dimensions['A'].width = 110
    for i in range(2, len(lines) + 2):
        summary[f'A{i}'].alignment = Alignment(wrap_text=True,
                                               vertical='top')
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--since', default='2026-09-01')
    ap.add_argument('--out', default='')
    args = ap.parse_args()

    out = args.out or os.path.join(
        'docs', 'exports',
        f'vertical_repair_{date.today().isoformat().replace("-", "_")}.xlsx')

    with app.app_context():
        print(f'database: {db.engine.url}')
        print(f'window:   Heavy Transport leads updated since {args.since}')
        print('MODE:     READ-ONLY — no lead is changed\n')

        data = va.survey(since=args.since)
        print(f'  examined       {data["examined"]:>6}')
        for title, key, _cols in SHEETS:
            print(f'  {title:<14} {len(data[key]):>6}')

        os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
        _write(out, data, args.since)
        print(f'\n  written to {out}')

        if data['changed_since']:
            print('\n  Changed since the rewrite — excluded from the '
                  'restore:')
            for row in data['changed_since'][:20]:
                print(f'    {row["lead_id"]:<8} {row["company"][:28]:<28} '
                      f'{row["was"]} -> Heavy Transport -> '
                      f'{row["changed_to"]} ({row["changed_by"]})')

        print('\nREAD-ONLY. Stage B applies sheet 1 only, on an explicit '
              'instruction.')


if __name__ == '__main__':
    main()
