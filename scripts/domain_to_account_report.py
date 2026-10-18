"""
The domains that became leads and matched no account.

READ-ONLY. Writes one .xlsx and changes nothing.

    .venv/bin/python scripts/domain_to_account_report.py --days 30

The match rate is 48 per cent, but most of what it misses is
newsletters and marketing, which should miss. The part that matters
is mail that became a lead and still matched nothing: a real
customer whose account the CRM cannot find, so the lead goes to a
fallback owner instead of to whoever knows them.

On 30 days that was 94 messages across 39 domains. Adding those
domains to their accounts does more for matching than any rule,
which is why this is a list for a person rather than something to
infer: `email_domains` decides who sees a customer's mail, and a
wrong row sends it to the wrong desk for ever.

Free mail is excluded. Nobody's account owns gmail.com.

The suggestion is the best name match the CRM can find, with how
confident it is. `confirm account` is blank, and blank is skipped.
"""
import argparse
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))

from app import Company, app, db                          # noqa: E402

COLUMNS = ('domain', 'messages', 'leads made', 'sample sender',
           'sample subject', 'suggested account', 'account id',
           'confidence', 'why', 'confirm account')

BECAME_A_LEAD = ('created', 'attached', 'reopened')


def _guess(domain, names):
    """(account, confidence, why) — the best name match, or nothing.

    The domain's own word against account names. `andritz.com` ->
    "andritz" -> an account called Andritz. Nothing clever, and
    nothing is written from it: a person confirms.
    """
    from app.services import account_match, company_match

    stem = re.sub(r'\.(com|net|org|in|co|io|de|cn|uk|us)(\.\w+)?$', '',
                  (domain or '').lower())
    stem = re.sub(r'^(mail|email|mailer|smtp|m|email\d*)\.', '', stem)
    stem = stem.replace('-', '').replace('.', '')
    if len(stem) < 3:
        return None, '', 'the domain has nothing to match on'

    exact = [c for c, key in names if key == stem]
    if len(exact) == 1:
        return exact[0], 'high', f'account name normalises to "{stem}"'
    if len(exact) > 1:
        return None, '', f'{len(exact)} accounts normalise to "{stem}"'

    starts = [c for c, key in names
              if key and (key.replace(' ', '').startswith(stem)
                          or stem.startswith(key.replace(' ', '')))]
    if len(starts) == 1:
        return starts[0], 'medium', (
            f'"{stem}" and "{starts[0].name}" start the same')
    if len(starts) > 1:
        return None, '', f'{len(starts)} accounts start with "{stem}"'
    return None, '', 'no account name resembles the domain'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--days', type=int, default=30)
    ap.add_argument('--out', default='')
    args = ap.parse_args()

    out = args.out or os.path.join(
        'docs', 'exports',
        f'domain_to_account_{date.today().isoformat().replace("-", "_")}'
        f'.xlsx')

    with app.app_context():
        from app.models.ingest_log import MailIngestLog
        from app.services import account_match, company_match

        print(f'database: {db.engine.url}')
        print('MODE:     READ-ONLY — nothing is written\n')

        since = datetime.utcnow() - timedelta(days=args.days)
        rows = (MailIngestLog.query
                .filter(MailIngestLog.received_at >= since).all())

        groups = defaultdict(lambda: {'messages': 0, 'leads': 0,
                                      'sender': '', 'subject': ''})
        for r in rows:
            addr = (r.from_addr or '').strip()
            if not addr or '@' not in addr:
                continue
            domain = account_match.domain_of(addr)
            if not domain or account_match.is_free_mail(domain):
                continue
            try:
                _acct, why = account_match.resolve(addr, count_hit=False)
            except Exception:
                db.session.rollback()
                continue
            if why != account_match.UNMAPPED:
                continue
            if r.outcome not in BECAME_A_LEAD:
                continue
            g = groups[domain]
            g['messages'] += 1
            g['leads'] += 1
            if not g['sender']:
                g['sender'] = addr
                g['subject'] = (r.subject or '')[:90]

        names = [(c, company_match.norm(c.name).replace(' ', ''))
                 for c in Company.query.filter(
                     Company.is_active.is_(True)).all()]

        table = []
        for domain, g in sorted(groups.items(),
                                key=lambda kv: -kv[1]['messages']):
            account, confidence, why = _guess(domain, names)
            table.append({
                'domain': domain, 'messages': g['messages'],
                'leads made': g['leads'], 'sample sender': g['sender'],
                'sample subject': g['subject'],
                'suggested account': account.name if account else '',
                'account id': account.id if account else '',
                'confidence': confidence, 'why': why,
                'confirm account': ''})

        print(f'  {len(table)} domain(s), '
              f'{sum(r["messages"] for r in table)} message(s)')
        for level, n in Counter(r['confidence'] or 'no suggestion'
                                for r in table).most_common():
            print(f'      {level:<16} {n:>4}')
        print()
        print(f'  {"domain":<34} {"msgs":>5}  {"suggested":<34} conf')
        for r in table:
            print(f'  {r["domain"][:34]:<34} {r["messages"]:>5}  '
                  f'{(r["suggested account"] or "—")[:34]:<34} '
                  f'{r["confidence"]}')

        _write(out, table)
        print(f'\n  written to {out}')
        print('\nREAD-ONLY. Nothing was written to any account.')


def _write(path, table):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Domains'
    ws['A1'] = (
        'Domains whose mail became a lead and matched no account. Put '
        'the account each belongs to in the last column — by name, '
        'exactly as the CRM holds it — and leave blank anything you '
        'are not sure of. Blanks are skipped. The suggestion is the '
        'best name match the CRM could find and is not an answer: '
        'email_domains decides who sees a customer\'s mail, and a '
        'wrong row sends it to the wrong desk for ever.')
    ws['A1'].font = Font(italic=True, size=9, color='6B6762')
    ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=len(COLUMNS))
    ws.row_dimensions[1].height = 58

    for col, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col, value=name)
        cell.fill = PatternFill('solid', fgColor='BC1D2F')
        cell.font = Font(color='FFFFFF', bold=True, size=10)
    for r, row in enumerate(table, start=3):
        for col, name in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=col, value=row.get(name, ''))

    widths = {'domain': 32, 'sample sender': 34, 'sample subject': 46,
              'suggested account': 34, 'why': 40, 'confirm account': 34,
              'messages': 10, 'leads made': 11, 'confidence': 12,
              'account id': 11}
    for col, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = widths[name]
    ws.freeze_panes = 'A3'
    if table:
        ws.auto_filter.ref = (f'A2:{get_column_letter(len(COLUMNS))}'
                              f'{len(table) + 2}')
    wb.save(path)


if __name__ == '__main__':
    main()
