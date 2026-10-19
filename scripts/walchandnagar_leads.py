"""
The 84 Walchandnagar leads, for a person to decide on.

    .venv/bin/python scripts/walchandnagar_leads.py

Read-only. Writes docs/exports/walchandnagar_leads_<date>.xlsx.

The client goes on the block register as BLOCKED, which stops new
work. It says nothing about the leads already on file, and those are
not a decision a script should take: some will be live quotes with
money attached and some will be two-year-old dead ends. Owner,
stage, created, last activity — enough to sort them in an afternoon.
"""
import os
import sys
from datetime import date, datetime

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _ROOT)

from app import Employee, Lead, app, db, STAGES_TERMINAL   # noqa: E402

MATCH = '%walchand%'


def main():
    with app.app_context():
        rows = (Lead.query.filter(Lead.company.ilike(MATCH))
                .order_by(Lead.stage.asc(), Lead.id.asc()).all())
        names = {e.emp_code: e.name for e in Employee.query.all()}
        from app.services import contact as _contact
        print(f'database: {db.engine.url}')
        print(f'{len(rows)} lead(s) whose company matches {MATCH!r}\n')

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Walchandnagar leads'
        ws.append(['Walchandnagar — leads already on file'])
        ws['A1'].font = Font(bold=True, size=14)
        ws.append([f'{len(rows)} lead(s), read-only snapshot '
                   f'{datetime.utcnow():%Y-%m-%d %H:%M} UTC. The client '
                   f'is BLOCKED for new work; these are not touched.'])
        ws.append([])
        cols = ['Lead', 'Company', 'Project', 'Stage', 'Open?', 'Owner',
                'Owner name', 'Created', 'Last activity', 'Days idle',
                'Value (INR)', 'Vertical', 'Decision (yours)']
        ws.append(cols)
        for cell in ws[4]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F2937')

        open_n = 0
        for l in rows:
            is_open = (l.stage or '') not in STAGES_TERMINAL
            open_n += 1 if is_open else 0
            last = max([d for d in (l.phone_call_date, l.intro_mail_date,
                                    l.meeting_date) if d] or [None]) \
                if any((l.phone_call_date, l.intro_mail_date,
                        l.meeting_date)) else None
            idle = ((date.today() - last).days if last else None)
            ws.append([
                l.id, l.company or '', l.project or '', l.stage or '',
                'open' if is_open else 'closed',
                l.assigned_to or '', names.get(l.assigned_to, ''),
                str(l.created_at)[:10] if l.created_at else '',
                str(last) if last else 'never',
                idle if idle is not None else '',
                float(l.quoted_amount_inr or l.estimated_value_inr or 0)
                or '',
                l.procam_vertical or '', ''])

        widths = (8, 34, 34, 20, 9, 12, 22, 12, 13, 10, 14, 20, 24)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f'A4:M{ws.max_row}'

        out = os.path.join(_ROOT, 'docs', 'exports',
                           f'walchandnagar_leads_'
                           f'{date.today().isoformat()}.xlsx')
        os.makedirs(os.path.dirname(out), exist_ok=True)
        wb.save(out)

        print(f'  open   {open_n}')
        print(f'  closed {len(rows) - open_n}')
        print(f'\n  {out}')
        print('  Nothing was written to the database.')


if __name__ == '__main__':
    main()
