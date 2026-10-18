"""
CRM-VERTICAL-MASTER-01 Stage A — settle the vertical list.

READ-ONLY. Writes one .xlsx and changes no lead, account or head.

Four questions, each a sheet:

  Values   every distinct vertical on a lead or an account, counted
           open and closed, with whether it is in the head master,
           the alias table, or neither — and a proposed official
           vertical marked `certain` or `needs Nilesh`.
  Heads    who each vertical belongs to, whether they are still
           active, and which verticals belong to nobody.
  Variants values that differ from another only by case, spacing or
           punctuation. A typo is not a vertical.
  About    the summary, and what Stage B would move.

`certain` means the CRM already knows: the value is a vertical in the
master, or the alias table maps it. Everything else is `needs
Nilesh`, because which vertical a lead belongs to is a business
decision and a script that guessed would be the same defect in
another costume.

    .venv/bin/python scripts/vertical_master_survey.py
"""
import argparse
import os
import re
import sys
from collections import defaultdict
from datetime import date

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))

from app import Company, Employee, Lead, app, db          # noqa: E402
from app.services import vertical_audit as va             # noqa: E402

BLANK = '(blank)'
#: Organisational groupings on the employee master, not services a
#: lead can be in.
NOT_A_SERVICE = ('All', 'Management')


def _norm(value):
    """A value stripped to its bones, for spotting variants."""
    return re.sub(r'[^a-z0-9]+', '', (value or '').lower())


def gather():
    from app.services import sales_rules as rules

    open_stages = set(rules.open_stages())
    rows = defaultdict(lambda: {'lead_open': 0, 'lead_closed': 0,
                                'accounts': 0})

    for vertical, stage, n in (
            db.session.query(Lead.procam_vertical, Lead.stage,
                             db.func.count(Lead.id))
            .filter(Lead.is_archived.isnot(True))
            .group_by(Lead.procam_vertical, Lead.stage).all()):
        key = (vertical or '').strip() or BLANK
        bucket = 'lead_open' if (stage or '') in open_stages else 'lead_closed'
        rows[key][bucket] += n

    try:
        for vertical, n in (db.session.query(Company.vertical,
                                             db.func.count(Company.id))
                            .group_by(Company.vertical).all()):
            key = (vertical or '').strip() or BLANK
            rows[key]['accounts'] += n
    except Exception:
        db.session.rollback()
    return rows


def heads():
    out = []
    try:
        from app.models.email_admin import VerticalHead
        masters = VerticalHead.query.order_by(VerticalHead.vertical).all()
    except Exception:
        db.session.rollback()
        return out
    by_code = {e.emp_code: e for e in Employee.query.all()}
    for row in masters:
        head = by_code.get((row.head_emp_code or '').strip().upper())
        deputy = by_code.get((row.deputy_emp_code or '').strip().upper())
        out.append({
            'vertical': row.vertical,
            'head code': row.head_emp_code or '',
            'head name': (head.name if head else
                          ('— nobody named —' if not row.head_emp_code
                           else 'NOT AN EMPLOYEE ON FILE')),
            'head active': ('' if head is None
                            else ('yes' if head.is_active else 'NO')),
            'deputy code': row.deputy_emp_code or '',
            'deputy name': (deputy.name if deputy else ''),
            'row active': 'yes' if row.active else 'no',
            'status': ('NO HEAD' if not (row.head_emp_code or '').strip()
                       else ('HEAD HAS LEFT' if head is not None
                             and not head.is_active
                             else ('HEAD NOT ON FILE' if head is None
                                   else 'ok'))),
        })
    return out


def propose(value, counts, masters, alias):
    """(proposed official vertical, certainty, note)."""
    if value == BLANK:
        return '', 'needs Nilesh', (
            f'{counts["lead_open"] + counts["lead_closed"]} leads with no '
            f'vertical at all. They match no head and appear in no '
            f'vertical report.')
    if value in alias:
        return alias[value], 'certain', (
            f'the alias table already maps this to {alias[value]}; '
            f'Stage B would move these leads')
    if value in masters and value not in NOT_A_SERVICE:
        return value, 'certain', 'a vertical in the head master'
    if value in NOT_A_SERVICE:
        return '', 'needs Nilesh', (
            'an organisational grouping on the employee master, not a '
            'service. A lead should probably not carry it.')
    return '', 'needs Nilesh', (
        'in use on records but in neither the head master nor the alias '
        'table, so it has no head and no mapping')


def write(path, values, head_rows, variants, choices, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    wb.remove(wb.active)
    red = PatternFill('solid', fgColor='BC1D2F')
    note_font = Font(italic=True, size=9, color='6B6762')

    def sheet(title, columns, rows, note, widths):
        ws = wb.create_sheet(title)
        ws['A1'] = note
        ws['A1'].font = note_font
        ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=max(4, len(columns)))
        ws.row_dimensions[1].height = 44
        for col, name in enumerate(columns, start=1):
            cell = ws.cell(row=2, column=col, value=name)
            cell.fill, cell.font = red, Font(color='FFFFFF', bold=True,
                                             size=10)
        for r, row in enumerate(rows, start=3):
            for col, name in enumerate(columns, start=1):
                ws.cell(row=r, column=col, value=row.get(name, ''))
        for col, name in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col)].width = \
                widths.get(name, 14)
        ws.freeze_panes = 'A3'
        if rows:
            ws.auto_filter.ref = (f'A2:{get_column_letter(len(columns))}'
                                  f'{len(rows) + 2}')
        return ws

    # ── About ──
    about = wb.create_sheet('About')
    about['A1'] = 'CRM-VERTICAL-MASTER-01 — the vertical list, for decision'
    about['A1'].font = Font(bold=True, size=13)
    for i, line in enumerate(summary, start=3):
        about[f'A{i}'] = line
        about[f'A{i}'].alignment = Alignment(wrap_text=True, vertical='top')
    about.column_dimensions['A'].width = 108

    # ── Values ──
    cols = ('value', 'lead_open', 'lead_closed', 'lead_total', 'accounts',
            'in head master', 'in alias table', 'proposed official vertical',
            'certainty', 'why', 'Nilesh decision')
    ws = sheet(
        'Values', cols, values,
        'Every distinct vertical on a lead or an account. "certain" means '
        'the CRM already knows — the value is in the head master, or the '
        'alias table maps it. Everything else needs your decision: put '
        'the official vertical in the last column. Leave it blank to '
        'leave the value alone.',
        {'value': 24, 'why': 56, 'proposed official vertical': 26,
         'certainty': 14, 'Nilesh decision': 22, 'in head master': 15,
         'in alias table': 20})

    lookup = wb.create_sheet('Verticals')
    lookup['A1'] = 'Valid verticals — do not edit'
    for i, v in enumerate(choices, start=2):
        lookup[f'A{i}'] = v
    lookup.sheet_state = 'hidden'

    answer = get_column_letter(len(cols))
    dv = DataValidation(type='list',
                        formula1=f'=Verticals!$A$2:$A${len(choices) + 1}',
                        allow_blank=True, showDropDown=False)
    dv.error = 'Pick one from the list, or leave blank to decide later.'
    dv.errorTitle = 'Not a vertical'
    ws.add_data_validation(dv)
    dv.add(f'{answer}3:{answer}{max(3, len(values) + 2)}')

    # ── Heads ──
    sheet('Heads',
          ('vertical', 'head code', 'head name', 'head active',
           'deputy code', 'deputy name', 'row active', 'status'),
          head_rows,
          'Who each vertical belongs to. "NO HEAD" means notifications '
          'for that vertical are copied to nobody above the owner. '
          'Management is the fallback for every unmapped vertical, so a '
          'gap there affects all of them.',
          {'vertical': 24, 'head name': 26, 'deputy name': 26,
           'status': 20})

    # ── Variants ──
    sheet('Variants',
          ('normalised', 'values that collapse to it', 'lead_total'),
          variants,
          'Values that differ from one another only by case, spacing or '
          'punctuation. A typo is not a vertical, and two spellings of '
          'one thing split its reporting in half.',
          {'normalised': 26, 'values that collapse to it': 60})

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='')
    args = ap.parse_args()

    out = args.out or os.path.join(
        'docs', 'exports',
        f'vertical_master_{date.today().isoformat().replace("-", "_")}.xlsx')

    with app.app_context():
        print(f'database: {db.engine.url}')
        print('MODE:     READ-ONLY — nothing is written\n')

        counts = gather()
        alias = {}
        masters = []
        try:
            from app.models.email_admin import VerticalAlias, VerticalHead
            alias = {a.alias: a.vertical for a in VerticalAlias.query.all()}
            masters = [v.vertical for v in VerticalHead.query.all()]
        except Exception:
            db.session.rollback()

        values = []
        for value in sorted(counts, key=lambda v: -(counts[v]['lead_open']
                                                    + counts[v]['lead_closed'])):
            c = counts[value]
            proposed, certainty, why = propose(value, c, masters, alias)
            values.append({
                'value': value,
                'lead_open': c['lead_open'],
                'lead_closed': c['lead_closed'],
                'lead_total': c['lead_open'] + c['lead_closed'],
                'accounts': c['accounts'],
                'in head master': 'yes' if value in masters else 'no',
                'in alias table': (f'yes -> {alias[value]}'
                                   if value in alias else 'no'),
                'proposed official vertical': proposed,
                'certainty': certainty,
                'why': why,
                'Nilesh decision': '',
            })

        groups = defaultdict(list)
        for value in counts:
            if value != BLANK:
                groups[_norm(value)].append(value)
        variants = [{'normalised': k,
                     'values that collapse to it': ', '.join(sorted(v)),
                     'lead_total': sum(counts[x]['lead_open']
                                       + counts[x]['lead_closed']
                                       for x in v)}
                    for k, v in sorted(groups.items()) if len(v) > 1]

        head_rows = heads()
        choices = va.acceptable_verticals(for_choice=True)

        certain = [v for v in values if v['certainty'] == 'certain']
        needs = [v for v in values if v['certainty'] == 'needs Nilesh']
        would_move = sum(v['lead_total'] for v in values
                         if v['value'] in alias)
        no_head = [h for h in head_rows if h['status'] != 'ok']

        summary = [
            f'Generated {date.today()} from the live database. Read-only: '
            f'nothing was changed.',
            '',
            f'{len(values)} distinct vertical values across leads and '
            f'accounts.',
            f'   {len(certain)} certain — the head master or the alias '
            f'table already settles them.',
            f'   {len(needs)} need your decision.',
            '',
            f'{would_move} leads carry a dead label the alias table '
            f'already maps. Stage B would move those without needing a '
            f'decision from you.',
            '',
            f'{len(head_rows)} rows in the vertical head master, '
            f'{len(no_head)} of them not in order:',
        ] + [f'   {h["vertical"]}: {h["status"]}' for h in no_head] + [
            '',
            'On the Values sheet, put the official vertical in the last '
            'column for every row marked "needs Nilesh". Leave it blank '
            'to leave that value alone. Stage B adds an alias row for '
            'each value you map elsewhere, sets the heads you specify, '
            'and moves the leads in one transaction with an audit row '
            'each.',
            '',
            'Leads still blank on the 296-row review files are never '
            'touched by Stage B.',
        ]

        print(f'  {len(values)} distinct values, {len(certain)} certain, '
              f'{len(needs)} need a decision')
        print(f'  {would_move} leads on a dead label would move')
        print(f'  {len(head_rows)} head rows, {len(no_head)} not in order\n')
        print(f'  {"value":<24} {"open":>6} {"closed":>7} {"acct":>6}  '
              f'{"certainty":<14} proposed')
        for v in values:
            print(f'  {v["value"]:<24} {v["lead_open"]:>6} '
                  f'{v["lead_closed"]:>7} {v["accounts"]:>6}  '
                  f'{v["certainty"]:<14} {v["proposed official vertical"]}')
        if variants:
            print('\n  variants that collapse together:')
            for v in variants:
                print(f'      {v["values that collapse to it"]}')
        print('\n  heads:')
        for h in head_rows:
            print(f'      {h["vertical"]:<22} {h["head code"]:<12} '
                  f'{h["head name"][:28]:<28} {h["status"]}')

        os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
        write(out, values, head_rows, variants, choices, summary)
        print(f'\n  written to {out}')
        print('\nREAD-ONLY. Nothing was changed.')


if __name__ == '__main__':
    main()
